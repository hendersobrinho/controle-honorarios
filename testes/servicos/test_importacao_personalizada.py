"""Importação personalizada: layout de colunas, modelo exportado e conferência.

Dados inventados. O foco é a regra pedida: cliente já cadastrado (PF ou PJ) não é
cadastrado de novo — só recebe o que a planilha traz de diferente — e honorário
que já existe não entra duas vezes.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from honorarios.dominio.enums import StatusPagamento, TipoPessoa
from honorarios.infra.modelos import Cliente, Contrato, Honorario
from honorarios.servicos import clientes as servico_clientes
from honorarios.servicos.erros import ErroDeValidacao
from honorarios.servicos.importacao import personalizada as imp
from testes.fabricas import CNPJ_FICTICIO, CPF_FICTICIO, cliente, honorario

CNPJ_NOVO = "11.444.777/0001-61"


def _layout(**extras) -> imp.Layout:
    campos = {
        "cliente_codigo": imp.Mapeamento("A"),
        "cliente_nome": imp.Mapeamento("B"),
        "cliente_documento": imp.Mapeamento("C"),
        "cliente_juros_dia": imp.Mapeamento("D"),
        "contato_telefone": imp.Mapeamento("E"),
        "honorario_tipo": imp.Mapeamento("", "Honorário mensal"),
        "honorario_competencia": imp.Mapeamento("F"),
        "honorario_vencimento": imp.Mapeamento("G"),
        "honorario_valor": imp.Mapeamento("H"),
    }
    campos.update(extras)
    return imp.Layout(nome="Teste", linha_inicial=2, campos=campos)


def _planilha(tmp_path, linhas, cabecalho=("Código", "Nome", "Doc")):
    livro = Workbook()
    aba = livro.active
    aba.append(list(cabecalho))
    for linha in linhas:
        aba.append(list(linha))
    arquivo = tmp_path / "entrada.xlsx"
    livro.save(arquivo)
    return arquivo


def _contar(sessao, modelo) -> int:
    return sessao.scalar(select(func.count(modelo.id)))


# ------------------------------------------------------------------- layout


def test_layout_exige_coluna_que_identifica_o_cliente():
    with pytest.raises(ErroDeValidacao, match="código, do CPF/CNPJ ou do nome"):
        imp.validar_layout(imp.Layout(campos={"honorario_valor": imp.Mapeamento("A")}))
    with pytest.raises(ErroDeValidacao, match="Coluna inválida"):
        imp.validar_layout(imp.Layout(campos={"cliente_nome": imp.Mapeamento("A1")}))


def test_salvar_listar_e_excluir_layout(sessao):
    imp.salvar_layout(sessao, _layout())
    layouts = imp.listar_layouts(sessao)
    assert list(layouts) == ["Teste"]
    assert layouts["Teste"].campos["cliente_nome"].coluna == "B"
    assert layouts["Teste"].campos["honorario_tipo"].padrao == "Honorário mensal"
    imp.excluir_layout(sessao, "Teste")
    assert imp.listar_layouts(sessao) == {}


def test_modelo_exportado_tem_cabecalho_nas_colunas_e_volta_como_layout(tmp_path):
    layout = _layout()
    arquivo = imp.exportar_modelo(layout, tmp_path / "modelo.xlsx")
    livro = load_workbook(arquivo)
    dados = livro.worksheets[0]
    assert dados["B1"].value == "Nome / razão social"
    assert dados["H1"].value == "Honorário: Valor"
    assert "Instruções" in livro.sheetnames
    assert livro[imp.ABA_LAYOUT].sheet_state == "hidden"

    de_volta = imp.ler_layout_do_modelo(arquivo)
    assert de_volta.para_dict() == layout.para_dict()


# --------------------------------------------------------- clientes e dedup


def test_cliente_ja_cadastrado_nao_duplica_e_recebe_o_que_falta(sessao, tmp_path):
    existente = cliente(sessao, codigo="900", documento=CNPJ_FICTICIO, juros_dia="1.01",
                        telefone="27999990000")
    arquivo = _planilha(tmp_path, [
        # mesmo CNPJ com outro código e telefone novo: é o mesmo cliente
        ["555", "EMPRESA EXEMPLO LTDA", "11222333000181", "1,50", "27988887777",
         "08/2026", date(2026, 9, 10), 800],
        # mesmo cliente de novo, outra competência
        ["555", "EMPRESA EXEMPLO LTDA", CNPJ_FICTICIO, "1,50", "", "09/2026",
         date(2026, 10, 10), 800],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    assert not any(linha.bloqueada for linha in previa.linhas)
    assert previa.linhas[0].acao_cliente.startswith("Atualizar 900")
    assert "juros/dia" in previa.linhas[0].acao_cliente
    assert any("código 900" in p.motivo for p in previa.todos_problemas())

    resultado = imp.importar(sessao, previa)
    sessao.commit()
    assert resultado.clientes_criados == 0
    assert resultado.clientes_atualizados == 1
    assert resultado.honorarios_criados == 2
    assert _contar(sessao, Cliente) == 1
    sessao.refresh(existente)
    assert existente.codigo == "900"
    assert servico_clientes.juros_atual(existente) == Decimal("1.5")
    assert {c.telefone for c in existente.contatos} == {"5527999990000", "5527988887777"}


def test_pf_reconhecida_pelo_cpf_e_pj_nova_criada(sessao, tmp_path):
    cliente(sessao, nome="MARIA EXEMPLO", codigo="10", documento=CPF_FICTICIO,
            tipo=TipoPessoa.PF)
    arquivo = _planilha(tmp_path, [
        ["", "Maria Exemplo", "52998224725", "", "", "08/2026", date(2026, 9, 5), 200],
        ["", "NOVA EMPRESA LTDA", CNPJ_NOVO, "0,80", "", "08/2026", date(2026, 9, 5), 300],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    resultado = imp.importar(sessao, previa)
    sessao.commit()
    assert resultado.clientes_criados == 1
    assert _contar(sessao, Cliente) == 2
    nova = servico_clientes.por_documento(sessao, CNPJ_NOVO)
    assert nova.tipo_pessoa is TipoPessoa.PJ
    assert nova.codigo == "11"   # próximo número livre
    maria = servico_clientes.por_codigo(sessao, "10")
    assert maria.nome == "MARIA EXEMPLO"   # só a caixa mudou: não é nome novo


def test_cliente_reconhecido_pelo_nome_quando_nao_ha_documento(sessao, tmp_path):
    cliente(sessao, nome="ÔMEGA SERVIÇOS LTDA", codigo="77", documento="")
    arquivo = _planilha(tmp_path, [
        ["", "omega servicos ltda", "", "", "", "08/2026", date(2026, 9, 5), 300],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    assert previa.linhas[0].acao_cliente.startswith("Já cadastrado (77)")
    imp.importar(sessao, previa)
    sessao.commit()
    assert _contar(sessao, Cliente) == 1


def test_codigo_de_outro_documento_trava_a_linha(sessao, tmp_path):
    cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    arquivo = _planilha(tmp_path, [
        ["900", "OUTRA EMPRESA", CNPJ_NOVO, "", "", "08/2026", date(2026, 9, 5), 300],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    assert previa.linhas[0].bloqueada
    resultado = imp.importar(sessao, previa)
    assert resultado.clientes_criados == 0 and resultado.honorarios_criados == 0


def test_cliente_novo_sem_nome_e_documento_invalido_travam(sessao, tmp_path):
    arquivo = _planilha(tmp_path, [
        ["123", "", "", "", "", "08/2026", date(2026, 9, 5), 300],
        ["124", "X LTDA", "11.111.111/1111-11", "", "", "08/2026", date(2026, 9, 5), 300],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    assert all(linha.bloqueada for linha in previa.linhas)
    motivos = " ".join(p.motivo for p in previa.todos_problemas())
    assert "não traz o nome" in motivos and "CPF/CNPJ inválido" in motivos


def test_sem_atualizar_existentes_so_lanca_o_novo(sessao, tmp_path):
    existente = cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    arquivo = _planilha(tmp_path, [
        ["900", "NOME TROCADO", CNPJ_FICTICIO, "", "", "08/2026", date(2026, 9, 5), 300],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    imp.conferir(sessao, previa, atualizar_existentes=False)
    assert "sem atualizar" in previa.linhas[0].acao_cliente
    resultado = imp.importar(sessao, previa, atualizar_existentes=False)
    sessao.commit()
    assert resultado.honorarios_criados == 1
    sessao.refresh(existente)
    assert existente.nome == "EMPRESA EXEMPLO LTDA"


# ------------------------------------------------------ honorários e contratos


def test_honorario_ja_lancado_nao_entra_de_novo(sessao, tmp_path):
    existente = cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    honorario(sessao, existente)   # 08/2026 já lançado
    arquivo = _planilha(tmp_path, [
        ["900", "", "", "", "", "08/2026", date(2026, 9, 8), 1507.80],
        ["900", "", "", "", "", "09/2026", date(2026, 10, 8), 1507.80],
        ["900", "", "", "", "", "09/2026", date(2026, 10, 8), 1507.80],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    acoes = [linha.acao_honorario for linha in previa.linhas]
    assert acoes[0].startswith("Já cadastrado")
    assert acoes[1].startswith("Novo")
    assert acoes[2].startswith("Repetido (linha 3)")
    resultado = imp.importar(sessao, previa)
    sessao.commit()
    assert resultado.honorarios_criados == 1
    assert _contar(sessao, Honorario) == 2


def test_honorario_pago_entra_congelado_com_data(sessao, tmp_path):
    cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    layout = _layout(honorario_status=imp.Mapeamento("I"), honorario_pago_em=imp.Mapeamento("J"))
    arquivo = _planilha(tmp_path, [
        ["900", "", "", "", "", "07/2026", date(2026, 8, 8), 500, "Pago", date(2026, 8, 7)],
    ])
    imp.importar(sessao, imp.ler(sessao, arquivo, layout))
    sessao.commit()
    pago = sessao.scalar(select(Honorario))
    assert pago.status is StatusPagamento.PAGO
    assert pago.saldo_aberto == 0 and pago.total_congelado == Decimal("500")
    assert pago.pago_em == date(2026, 8, 7)


def test_honorario_sem_vencimento_e_tipo_desconhecido_travam(sessao, tmp_path):
    layout = _layout(honorario_tipo=imp.Mapeamento("I"))
    arquivo = _planilha(tmp_path, [
        ["1", "A LTDA", "", "", "", "08/2026", None, 300, "Honorário mensal"],
        ["2", "B LTDA", "", "", "", "08/2026", date(2026, 9, 5), 300, "Inexistente"],
    ])
    previa = imp.ler(sessao, arquivo, layout)
    assert all(linha.bloqueada for linha in previa.linhas)


def test_contrato_criado_e_depois_reajustado(sessao, tmp_path):
    layout = imp.Layout(nome="Contratos", campos={
        "cliente_codigo": imp.Mapeamento("A"),
        "cliente_nome": imp.Mapeamento("B"),
        "contrato_tipo": imp.Mapeamento("", "Honorário mensal"),
        "contrato_valor": imp.Mapeamento("C"),
        "contrato_dia_vencimento": imp.Mapeamento("D"),
    })
    imp.importar(sessao, imp.ler(sessao, _planilha(tmp_path, [["50", "C LTDA", 900, 15]]),
                                 layout))
    sessao.commit()
    contrato = sessao.scalar(select(Contrato))
    assert contrato.valor == Decimal("900") and contrato.dia_vencimento == 15

    previa = imp.ler(sessao, _planilha(tmp_path, [["50", "C LTDA", 990, 20]]), layout)
    assert previa.linhas[0].acao_contrato.startswith("Reajuste")
    resultado = imp.importar(sessao, previa)
    sessao.commit()
    assert resultado.contratos_atualizados == 1
    assert _contar(sessao, Contrato) == 1
    sessao.refresh(contrato)
    assert contrato.valor == Decimal("990") and contrato.dia_vencimento == 20
    assert contrato.historico[-1].valor_anterior == "900.00"


def test_pf_vinculada_a_empresa_nova_da_mesma_planilha(sessao, tmp_path):
    layout = imp.Layout(nome="Vínculo", campos={
        "cliente_codigo": imp.Mapeamento("A"),
        "cliente_nome": imp.Mapeamento("B"),
        "cliente_documento": imp.Mapeamento("C"),
        "empresa_vinculo": imp.Mapeamento("D"),
        "suspensao_motivo": imp.Mapeamento("E"),
    })
    arquivo = _planilha(tmp_path, [
        ["60", "SÓCIO EXEMPLO", CPF_FICTICIO, "61", "Em negociação"],
        ["61", "EMPRESA DO SÓCIO LTDA", CNPJ_NOVO, "", ""],
    ])
    previa = imp.ler(sessao, arquivo, layout)
    assert not any(linha.bloqueada for linha in previa.linhas)
    imp.importar(sessao, previa)
    sessao.commit()
    socio = servico_clientes.por_codigo(sessao, "60")
    empresa = servico_clientes.por_codigo(sessao, "61")
    assert socio.empresa_id == empresa.id
    assert socio.cobranca_suspensa and socio.motivo_suspensao == "Em negociação"


def test_valor_fixo_do_layout_vale_quando_a_celula_esta_vazia(sessao, tmp_path):
    layout = _layout(cliente_juros_dia=imp.Mapeamento("D", "0,55"))
    arquivo = _planilha(tmp_path, [
        ["70", "D LTDA", "", "", "", "08/2026", date(2026, 9, 5), 300],
    ])
    imp.importar(sessao, imp.ler(sessao, arquivo, layout))
    sessao.commit()
    novo = servico_clientes.por_codigo(sessao, "70")
    assert servico_clientes.juros_atual(novo) == Decimal("0.55")


# ------------------------------------------------ achados da revisão de código


def test_codigo_gerado_nao_rouba_o_de_uma_linha_mais_abaixo(sessao, tmp_path):
    cliente(sessao, codigo="500", documento=CNPJ_FICTICIO)
    arquivo = _planilha(tmp_path, [
        ["", "SEM CODIGO LTDA", "", "", "", "08/2026", date(2026, 9, 5), 100],
        ["501", "COM CODIGO LTDA", "", "", "", "08/2026", date(2026, 9, 5), 100],
    ])
    resultado = imp.importar(sessao, imp.ler(sessao, arquivo, _layout()))
    assert resultado.clientes_criados == 2 and not resultado.falhas
    assert servico_clientes.por_codigo(sessao, "501").nome == "COM CODIGO LTDA"
    assert servico_clientes.por_codigo(sessao, "502").nome == "SEM CODIGO LTDA"


def test_honorario_cancelado_nao_duplica_ao_reimportar(sessao, tmp_path):
    cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    layout = _layout(honorario_status=imp.Mapeamento("I"))
    arquivo = _planilha(tmp_path, [
        ["900", "", "", "", "", "08/2026", date(2026, 9, 8), 500, "Cancelado"],
    ])
    imp.importar(sessao, imp.ler(sessao, arquivo, layout))
    previa = imp.ler(sessao, arquivo, layout)
    assert previa.linhas[0].acao_honorario.startswith("Já cadastrado")
    imp.importar(sessao, previa)
    assert _contar(sessao, Honorario) == 1


def test_contato_so_com_email_e_gravado(sessao, tmp_path):
    existente = cliente(sessao, codigo="900", documento=CNPJ_FICTICIO)
    layout = imp.Layout(campos={
        "cliente_codigo": imp.Mapeamento("A"), "contato_email": imp.Mapeamento("B"),
    })
    arquivo = _planilha(tmp_path, [["900", "financeiro@exemplo.com"]])
    imp.importar(sessao, imp.ler(sessao, arquivo, layout))
    sessao.commit()
    sessao.refresh(existente)
    assert "financeiro@exemplo.com" in {c.email for c in existente.contatos}
    # reimportar não acha mais nada para atualizar
    previa = imp.ler(sessao, arquivo, layout)
    assert previa.linhas[0].acao_cliente.startswith("Já cadastrado")


def test_reconhecido_so_pelo_codigo_nao_troca_o_nome(sessao, tmp_path):
    existente = cliente(sessao, nome="BELTRANO", codigo="1", documento="")
    arquivo = _planilha(tmp_path, [
        ["1", "FULANO", "", "", "", "08/2026", date(2026, 9, 5), 100],
    ])
    previa = imp.ler(sessao, arquivo, _layout())
    assert any("Confira se é o mesmo cliente" in p.motivo for p in previa.todos_problemas())
    imp.importar(sessao, previa)
    sessao.refresh(existente)
    assert existente.nome == "BELTRANO"


def test_sem_atualizar_existentes_contrato_fica_como_esta(sessao, tmp_path):
    layout = imp.Layout(campos={
        "cliente_codigo": imp.Mapeamento("A"), "cliente_nome": imp.Mapeamento("B"),
        "contrato_tipo": imp.Mapeamento("", "Honorário mensal"),
        "contrato_valor": imp.Mapeamento("C"),
    })
    imp.importar(sessao, imp.ler(sessao, _planilha(tmp_path, [["50", "C LTDA", 900]]), layout))
    previa = imp.ler(sessao, _planilha(tmp_path, [["50", "C LTDA", 990]]), layout,
                     atualizar_existentes=False)
    assert previa.linhas[0].acao_contrato == "Já cadastrado, sem atualizar"
    imp.importar(sessao, previa, atualizar_existentes=False)
    assert sessao.scalar(select(Contrato)).valor == Decimal("900")
