"""Importação: planilha sintética que reproduz a sujeira da planilha real.

Nenhum dado real de cliente entra aqui — nomes, CPF/CNPJ e valores são
inventados, mas os defeitos são os mesmos encontrados na planilha em uso:
competência em quatro formatos, linha sem vencimento, linha sem valor, status
desconhecido, juros ao dia divergente e cobranças como texto na observação.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from honorarios.dominio.enums import ResultadoCobranca, StatusPagamento
from honorarios.servicos import honorarios as servico_honorarios
from honorarios.servicos.honorarios import Filtro
from honorarios.servicos.importacao import executor, leitor

CABECALHO_BASE = [
    "Competência", "Código Cliente", "Cliente", "CPF/CNPJ", "Status Pagamento",
    "Vencimento", "Valor do Boleto", "Juros ao Dia", "Dias em Atraso", "Juros até Hoje",
    "Total Atualizado", "Total Débitos em Aberto", "Observações", "Tipo",
    "Tipo de Faturamento", "Forma dos Honorários",
]
CABECALHO_IRPF = [
    "Tipo", "Tipo de Faturamento", "Competência", "Código Cliente", "Cliente", "CPF/CNPJ",
    "Vencimento", "Valor do Boleto", "Juros ao Dia", "Forma dos Honorários", "Boleto ou Pix",
    "Status Pagamento", "Dias em Atraso", "Juros até Hoje", "Total Atualizado",
    "Há Parcela Anterior em Aberto?", "Valor em Aberto Anterior no Boleto", "Observações",
    "Arquivo Origem",
]

CNPJ = "11.222.333/0001-81"
CPF = "529.982.247-25"
CPF_2 = "398.117.700-04"


@pytest.fixture
def planilha(tmp_path):
    livro = Workbook()
    base = livro.active
    base.title = "Controle Financeiro"
    base.append(["BOLETOS EM ABERTOS"])
    base.append(CABECALHO_BASE)
    linhas = [
        # competência texto, em aberto, com cobranças na observação
        ["08/2026", 900, "ALFA COMERCIO LTDA", CNPJ, "Em aberto", datetime(2026, 9, 8),
         1000.00, 1.01, None, None, None, None,
         "17/09/2026 - Feito cobrança.\n22/09/2026 - Disse que vai pagar dia 30/09/2026.",
         "PJ", "Mensal", "Onvio"],
        # competência como data
        [datetime(2026, 7, 1), 900, "ALFA COMERCIO LTDA", CNPJ, "Pago", datetime(2026, 8, 8),
         1000.00, 1.01, None, None, None, None, None, "PJ", "Mensal", "Onvio"],
        # competência como serial do Excel (46174 = 06/2026)
        [46174, 901, "BETA SERVICOS ME", "", "Em aberto", datetime(2026, 7, 6),
         500.00, 0.34, None, None, None, None, None, "PJ", "Mensal", "E-mail"],
        # competência com o nome do serviço
        ["ATA DE SOCIOS", 901, "BETA SERVICOS ME", "", "Pago", datetime(2026, 1, 20),
         600.00, None, None, None, None, None, None, "PJ", "Ata de Socios", "Malote"],
        # sem vencimento -> erro
        ["05/2026", 902, "GAMA PARTICIPACOES LTDA", "", "Pago", None,
         480.00, 0.32, None, None, None, None, None, "PJ", "Mensal", "E-mail"],
        # valor zerado -> erro
        [None, 903, "DELTA ASSESSORIA LTDA", "", "Pago", datetime(2026, 6, 5),
         0, 0, None, None, None, None, None, "PJ", "Mensal", "E-mail"],
        # status desconhecido
        ["08/2026", 904, "EPSILON LTDA", "", "Aguardando", datetime(2026, 9, 10),
         300.00, 0.20, None, None, None, None, None, "PJ", "Mensal", "Onvio"],
        # doméstica com status de envio
        ["08/2026", 905, "MARIA DA SILVA EXEMPLO", CPF, "Enviado para Cliente",
         datetime(2026, 9, 5), 200.00, 0.13, None, None, None, None, None,
         "Doméstica", "Mensal", "WhatsApp do cliente"],
        # trimestral
        ["06/2026", 906, "ZETA HOLDING LTDA", "", "Em aberto", datetime(2026, 7, 6),
         660.00, 0.44, None, None, None, None, None, "PJ", "Trimestral", "E-mail"],
        # mesmo cliente do 901 com juros divergente
        ["09/2026", 901, "BETA SERVICOS ME", "", "Em aberto", datetime(2026, 10, 6),
         500.00, 2.16, None, None, None, None, None, "PJ", "Mensal", "E-mail"],
        # CPF inválido
        ["08/2026", 907, "ETA CONSULTORIA LTDA", "11.111.111/1111-11", "Em aberto",
         datetime(2026, 9, 12), 150.00, 0.10, None, None, None, None, None,
         "PJ", "Mensal", "Onvio"],
    ]
    for linha in linhas:
        base.append(linha)

    irpf = livro.create_sheet("Controle de IRPF")
    irpf.append(["Controle Financeiro de Honorários - 2026"])
    irpf.append(CABECALHO_IRPF)
    irpf.append(
        ["Imposto de Renda", "Imposto de Renda", "IRPF 2026", 2100, "JOAO EXEMPLO SOUZA",
         CPF_2, datetime(2026, 6, 25), 550.00, 0.37, "E-mail", "Boleto", "Pago",
         0, 0, 0, "Não", 0, None, None]
    )
    irpf.append(
        ["Imposto de Renda", "Imposto de Renda", "IRPF 2026", 900, "ALFA COMERCIO LTDA",
         CNPJ, datetime(2026, 6, 25), 700.00, 0.47, "E-mail", "Pix", "Em aberto",
         0, 0, 0, "Não", 0, "30/06/2026 - Feito cobrança.", None]
    )

    recall = livro.create_sheet("Recall Inadimplentes")
    recall.append(["Recall de Clientes Inadimplentes"])
    recall.append(["Ano", 2026])
    recall.append([])
    recall.append(["Código Cliente", "Cliente", "Tipo", "Qtd. Em Aberto"])
    recall.append(["900", "ALFA COMERCIO LTDA", "PJ", 1])
    recall.append(["950", "TETA LTDA (SEM LANCAMENTO)", "PJ", 0])

    destino = tmp_path / "planilha-exemplo.xlsx"
    livro.save(destino)
    return destino


# ------------------------------------------------------------------- leitura


def test_leitura_normaliza_as_quatro_formas_de_competencia(planilha):
    leitura = leitor.ler(planilha)
    rotulos = [linha.competencia.rotulo() for linha in leitura.linhas]
    assert "08/2026" in rotulos
    assert "07/2026" in rotulos       # veio como data
    assert "06/2026" in rotulos       # veio como serial do Excel
    assert "2026" in rotulos          # IRPF


def test_competencia_com_nome_de_servico_vira_tipo(planilha):
    leitura = leitor.ler(planilha)
    ata = [linha for linha in leitura.linhas if linha.tipo_honorario == "Ata de Sócios"]
    assert len(ata) == 1
    assert ata[0].competencia.vazia
    assert any(p.campo == "Competência" for p in ata[0].problemas)


def test_linhas_impossiveis_ficam_bloqueadas_com_motivo(planilha):
    leitura = leitor.ler(planilha)
    bloqueadas = [linha for linha in leitura.linhas if linha.bloqueada]
    motivos = {p.campo for linha in bloqueadas for p in linha.problemas if p.bloqueia}
    assert motivos == {"Vencimento", "Valor do Boleto"}
    assert len(bloqueadas) == 2


def test_cobrancas_saem_do_texto_das_observacoes(planilha):
    leitura = leitor.ler(planilha)
    com_cobranca = [linha for linha in leitura.linhas if linha.cobrancas]
    assert sum(len(linha.cobrancas) for linha in com_cobranca) == 3
    primeira = com_cobranca[0].cobrancas[0]
    assert primeira.quando == date(2026, 9, 17)
    assert primeira.texto == "Feito cobrança."


def test_status_desconhecido_vira_em_aberto_com_aviso(planilha):
    leitura = leitor.ler(planilha)
    linha = next(item for item in leitura.linhas if item.nome == "EPSILON LTDA")
    assert linha.status is StatusPagamento.EM_ABERTO
    assert any(p.campo == "Status Pagamento" for p in linha.problemas)


def test_enviado_para_cliente_e_status_de_envio(planilha):
    leitura = leitor.ler(planilha)
    linha = next(item for item in leitura.linhas if item.nome == "MARIA DA SILVA EXEMPLO")
    assert linha.status is StatusPagamento.EM_ABERTO
    assert linha.situacao_envio.name == "ENVIADO"
    assert linha.subtipo.name == "DOMESTICA"


def test_documento_invalido_e_descartado_com_aviso(planilha):
    leitura = leitor.ler(planilha)
    linha = next(item for item in leitura.linhas if item.nome == "ETA CONSULTORIA LTDA")
    assert linha.documento == ""
    assert any(p.campo == "CPF/CNPJ" for p in leitura.problemas)


def test_juros_divergente_gera_aviso_e_sugere_o_maior(planilha):
    leitura = leitor.ler(planilha)
    aviso = [
        p for p in leitura.problemas
        if p.campo == "Juros ao Dia" and "divergentes" in p.motivo
    ]
    # BETA (0,34 e 2,16 na base) e ALFA (1,01 na base, 0,47 no IRPF)
    assert len(aviso) == 2
    beta = next(c for c in leitura.clientes.values() if c.nome == "BETA SERVICOS ME")
    assert beta.juros_dia == Decimal("2.1600")


def test_tipo_e_corrigido_pelo_documento(planilha):
    """A coluna Tipo do IRPF diz "Imposto de Renda" mesmo em linha de empresa."""
    leitura = leitor.ler(planilha)
    alfa = next(c for c in leitura.clientes.values() if c.nome == "ALFA COMERCIO LTDA")
    assert alfa.tipo_pessoa.value == "PJ"
    assert any(p.campo == "Tipo" for p in leitura.problemas)


def test_recall_semeia_cliente_sem_lancamento(planilha):
    leitura = leitor.ler(planilha)
    nomes = {c.nome for c in leitura.clientes.values()}
    assert "TETA LTDA (SEM LANCAMENTO)" in nomes


def test_mesmo_cliente_nas_tres_abas_nao_duplica(planilha):
    leitura = leitor.ler(planilha)
    alfas = [c for c in leitura.clientes.values() if c.nome == "ALFA COMERCIO LTDA"]
    assert len(alfas) == 1


# ------------------------------------------------------------------ gravação


def test_importacao_grava_clientes_honorarios_e_cobrancas(sessao, planilha):
    leitura = leitor.ler(planilha)
    resultado = executor.importar(sessao, leitura)
    sessao.commit()

    assert resultado.clientes_criados == len(leitura.clientes)
    assert resultado.honorarios_criados == len(leitura.importaveis)
    assert resultado.cobrancas_criadas == 3
    assert resultado.honorarios_ignorados == 2

    pendentes = servico_honorarios.pendentes(
        sessao, Filtro(data_referencia=date(2026, 9, 29))
    )
    assert len(pendentes) == 8
    cobradas = [linha for linha in pendentes if linha.cobrancas]
    assert cobradas and cobradas[0].honorario.cobrancas[0].origem_texto_planilha


def test_pago_entra_congelado_e_com_explicacao(sessao, planilha):
    leitura = leitor.ler(planilha)
    executor.importar(sessao, leitura)
    sessao.commit()
    pagos = servico_honorarios.pagos(sessao, Filtro(data_referencia=date(2026, 9, 29)))
    assert pagos
    honorario = pagos[0].honorario
    assert honorario.total_congelado == honorario.valor_original
    assert honorario.saldo_aberto == Decimal("0.00")
    assert "não guardava data nem valor" in honorario.observacoes
    assert pagos[0].atualizacao.congelado


def test_cobranca_importada_infere_resultado_do_texto(sessao, planilha):
    leitura = leitor.ler(planilha)
    executor.importar(sessao, leitura)
    sessao.commit()
    from honorarios.infra.modelos import Cobranca

    resultados = {c.resultado for c in sessao.query(Cobranca).all()}
    assert ResultadoCobranca.PROMETEU_PAGAR in resultados
    assert ResultadoCobranca.SEM_RESPOSTA in resultados


def test_importar_duas_vezes_nao_duplica(sessao, planilha):
    leitura = leitor.ler(planilha)
    executor.importar(sessao, leitura)
    sessao.commit()
    segunda = executor.importar(sessao, leitor.ler(planilha))
    sessao.commit()
    assert segunda.clientes_criados == 0
    assert segunda.honorarios_criados == 0
    assert segunda.honorarios_ignorados >= len(leitura.importaveis)


def test_relatorio_de_erros_fica_gravado(sessao, planilha):
    leitura = leitor.ler(planilha)
    resultado = executor.importar(sessao, leitura)
    sessao.commit()
    from honorarios.infra.modelos import ImportacaoErro

    erros = sessao.query(ImportacaoErro).filter_by(importacao_id=resultado.importacao_id).all()
    assert any(e.gravidade == "ERRO" for e in erros)
    assert any(e.gravidade == "AVISO" for e in erros)


def test_opcao_de_nao_importar_pagos(sessao, planilha):
    leitura = leitor.ler(planilha)
    resultado = executor.importar(sessao, leitura, importar_pagos=False)
    sessao.commit()
    assert servico_honorarios.pagos(sessao, Filtro()) == []
    assert resultado.honorarios_criados < len(leitura.importaveis)
