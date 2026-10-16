"""Importação personalizada: o usuário diz em que coluna do Excel está cada dado.

O layout (nome, aba, linha onde começam os dados e, para cada informação do
cadastro, a letra da coluna e um valor fixo opcional) é guardado nas
configurações e pode ser exportado como planilha-modelo para preenchimento.

A leitura não grava nada: monta a prévia e já confere o que existe no banco.
Cliente é reconhecido pelo CPF/CNPJ, depois pelo código e, por último, pelo nome
exato — nunca é cadastrado de novo; dele só se atualiza o que a planilha traz de
diferente. Honorário que já existe (mesma competência, ou mesmo vencimento e
valor quando não há competência) não entra duas vezes; contrato do mesmo tipo é
atualizado, e mudança de valor vira reajuste com histórico.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field, replace
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ...dominio.enums import (
    ExigeCompetencia,
    MeioCobranca,
    Periodicidade,
    SituacaoCliente,
    SituacaoEnvio,
    StatusPagamento,
    SubtipoCliente,
    TipoPessoa,
)
from ...dominio.valores import (
    ZERO,
    Competencia,
    formatar_moeda,
    hoje,
    normalizar_busca,
    normalizar_telefone,
    so_digitos,
    validar_documento,
)
from ...infra.modelos import (
    Cliente,
    Contrato,
    FormaEnvio,
    Honorario,
    Importacao,
    ImportacaoErro,
    TipoHonorario,
)
from .. import clientes as servico_clientes
from .. import configuracoes
from .. import contratos as servico_contratos
from .. import honorarios as servico_honorarios
from ..erros import ErroDeNegocio, ErroDeValidacao
from .leitor import MAPA_STATUS, Problema, _data, _decimal, _juros, _texto

CHAVE_CONFIGURACAO = "importacao_layouts"
ABA_LAYOUT = "_layout"
MOTIVO = "Importação personalizada"

# --------------------------------------------------------------------- campos

GRUPO_CLIENTE = "Cliente"
GRUPO_CONTATO = "Contato"
GRUPO_VINCULO = "Vínculo e suspensão"
GRUPO_CONTRATO = "Contrato (honorário recorrente)"
GRUPO_HONORARIO = "Honorário (lançamento)"
GRUPOS = (GRUPO_CLIENTE, GRUPO_CONTATO, GRUPO_VINCULO, GRUPO_CONTRATO, GRUPO_HONORARIO)


@dataclass(frozen=True)
class Campo:
    chave: str
    rotulo: str
    grupo: str
    # texto, documento, dinheiro, taxa, data, inteiro, competencia, sim_nao,
    # telefone, pessoa, subtipo, situacao, periodicidade, meio, status, envio,
    # tipo_honorario, forma_envio
    tipo: str = "texto"
    dica: str = ""


CAMPOS: tuple[Campo, ...] = (
    Campo("cliente_codigo", "Código", GRUPO_CLIENTE,
          dica="Sem código, o cliente novo recebe o próximo número livre"),
    Campo("cliente_nome", "Nome / razão social", GRUPO_CLIENTE,
          dica="Obrigatório para cadastrar cliente novo"),
    Campo("cliente_documento", "CPF/CNPJ", GRUPO_CLIENTE, "documento",
          "Com ou sem pontuação; é o que reconhece o cliente já cadastrado"),
    Campo("cliente_tipo_pessoa", "Tipo de pessoa", GRUPO_CLIENTE, "pessoa",
          "PF ou PJ; vazio, vale o que o CPF/CNPJ indicar"),
    Campo("cliente_subtipo", "Subtipo", GRUPO_CLIENTE, "subtipo", "Doméstica ou vazio"),
    Campo("cliente_situacao", "Situação", GRUPO_CLIENTE, "situacao",
          "Ativo, Inativo ou Suspenso"),
    Campo("cliente_juros_dia", "Juros por dia (R$)", GRUPO_CLIENTE, "taxa",
          "Ex.: 1,01. Mudança vira nova vigência a partir de hoje"),
    Campo("cliente_observacoes", "Observações do cliente", GRUPO_CLIENTE,
          dica="Acrescentadas às que o cliente já tem"),
    Campo("contato_nome", "Nome do contato", GRUPO_CONTATO),
    Campo("contato_tratamento", "Como chamar (tratamento)", GRUPO_CONTATO,
          dica="O nome usado na saudação da mensagem"),
    Campo("contato_telefone", "Telefone / WhatsApp", GRUPO_CONTATO, "telefone",
          "Com DDD; contato com o mesmo telefone não é repetido"),
    Campo("contato_email", "E-mail", GRUPO_CONTATO),
    Campo("empresa_vinculo", "Empresa que paga por ele", GRUPO_VINCULO,
          dica="Código ou CNPJ da empresa (só para pessoa física)"),
    Campo("suspensao_motivo", "Suspender cobrança: motivo", GRUPO_VINCULO,
          dica="Preenchido, a cobrança do cliente fica suspensa"),
    Campo("suspensao_ate", "Suspensa até", GRUPO_VINCULO, "data",
          "Data de retomada; vazio, até retomar na mão"),
    Campo("contrato_tipo", "Tipo de honorário", GRUPO_CONTRATO, "tipo_honorario",
          "Nome do tipo como está no cadastro"),
    Campo("contrato_valor", "Valor do contrato", GRUPO_CONTRATO, "dinheiro",
          "Preenchido, cria (ou atualiza) o contrato do cliente"),
    Campo("contrato_periodicidade", "Periodicidade", GRUPO_CONTRATO, "periodicidade",
          "Mensal, Trimestral ou Anual; vazio, a do tipo"),
    Campo("contrato_dia_vencimento", "Dia do vencimento", GRUPO_CONTRATO, "inteiro",
          "1 a 31; vazio, dia 10"),
    Campo("contrato_meses_vencimento", "Vence quantos meses depois", GRUPO_CONTRATO,
          "inteiro", "Ex.: 1 = competência 08 vence em 09"),
    Campo("contrato_meses_referencia", "Meses de referência", GRUPO_CONTRATO,
          dica="Para trimestral/anual. Ex.: 3,6,9,12"),
    Campo("contrato_forma_envio", "Forma de envio", GRUPO_CONTRATO, "forma_envio"),
    Campo("contrato_meio", "Boleto ou Pix", GRUPO_CONTRATO, "meio"),
    Campo("contrato_vigencia_inicio", "Início da vigência", GRUPO_CONTRATO, "data",
          "Vazio, hoje"),
    Campo("contrato_vigencia_fim", "Fim da vigência", GRUPO_CONTRATO, "data"),
    Campo("contrato_suspenso", "Contrato suspenso", GRUPO_CONTRATO, "sim_nao", "Sim ou Não"),
    Campo("contrato_observacoes", "Observações do contrato", GRUPO_CONTRATO),
    Campo("honorario_tipo", "Tipo de honorário", GRUPO_HONORARIO, "tipo_honorario",
          "Nome do tipo como está no cadastro"),
    Campo("honorario_competencia", "Competência", GRUPO_HONORARIO, "competencia",
          "08/2026, 2026-08, data ou só o ano (IR)"),
    Campo("honorario_descricao", "Descrição", GRUPO_HONORARIO,
          dica="Obrigatória para tipo sem competência (avulso)"),
    Campo("honorario_valor", "Valor do honorário", GRUPO_HONORARIO, "dinheiro",
          "Preenchido, lança o honorário"),
    Campo("honorario_vencimento", "Vencimento", GRUPO_HONORARIO, "data",
          "Obrigatório para lançar o honorário"),
    Campo("honorario_status", "Status", GRUPO_HONORARIO, "status",
          "Em aberto, Pago, Cancelado ou Renegociado; vazio, em aberto"),
    Campo("honorario_pago_em", "Pago em", GRUPO_HONORARIO, "data",
          "Data do pagamento, quando o status é Pago"),
    Campo("honorario_meio", "Boleto ou Pix", GRUPO_HONORARIO, "meio"),
    Campo("honorario_forma_envio", "Forma de envio", GRUPO_HONORARIO, "forma_envio"),
    Campo("honorario_situacao_envio", "Enviado ao cliente", GRUPO_HONORARIO, "envio",
          "Enviado ou Não enviado"),
    Campo("honorario_numero_boleto", "Número do boleto", GRUPO_HONORARIO),
    Campo("honorario_observacoes", "Observações do honorário", GRUPO_HONORARIO),
)
POR_CHAVE = {campo.chave: campo for campo in CAMPOS}


def campos_do_grupo(grupo: str) -> list[Campo]:
    return [campo for campo in CAMPOS if campo.grupo == grupo]


# ---------------------------------------------------------------------- layout


@dataclass
class Mapeamento:
    coluna: str = ""
    padrao: str = ""

    @property
    def vazio(self) -> bool:
        return not self.coluna and not self.padrao


@dataclass
class Layout:
    nome: str = ""
    aba: str = ""
    linha_inicial: int = 2
    campos: dict[str, Mapeamento] = field(default_factory=dict)

    def usados(self) -> dict[str, Mapeamento]:
        return {
            chave: mapa for chave, mapa in self.campos.items()
            if chave in POR_CHAVE and not mapa.vazio
        }

    def usa(self, chave: str) -> bool:
        mapa = self.campos.get(chave)
        return mapa is not None and not mapa.vazio

    def para_dict(self) -> dict:
        return {
            "nome": self.nome,
            "aba": self.aba,
            "linha_inicial": self.linha_inicial,
            "campos": {
                chave: {"coluna": mapa.coluna, "padrao": mapa.padrao}
                for chave, mapa in self.usados().items()
            },
        }

    @classmethod
    def de_dict(cls, dados: dict) -> Layout:
        campos = {}
        for chave, mapa in (dados.get("campos") or {}).items():
            if chave not in POR_CHAVE or not isinstance(mapa, dict):
                continue
            campos[chave] = Mapeamento(
                coluna=str(mapa.get("coluna") or "").strip().upper(),
                padrao=str(mapa.get("padrao") or "").strip(),
            )
        try:
            linha = int(dados.get("linha_inicial") or 2)
        except (TypeError, ValueError):
            linha = 2
        return cls(
            nome=str(dados.get("nome") or "").strip(),
            aba=str(dados.get("aba") or "").strip(),
            linha_inicial=max(1, linha),
            campos=campos,
        )


def indice_coluna(letra: str) -> int:
    """Letra do Excel (A, B… XFD) para índice começando em 1."""
    texto = (letra or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", texto):
        raise ErroDeValidacao(f"Coluna inválida: {letra!r}. Use letras, como A, B ou AA.")
    indice = column_index_from_string(texto)
    if indice > 16384:
        raise ErroDeValidacao(f"A coluna {texto} passa do limite do Excel (XFD).")
    return indice


def validar_layout(layout: Layout) -> None:
    usados = layout.usados()
    if not any(mapa.coluna for mapa in usados.values()):
        raise ErroDeValidacao("Informe a coluna de pelo menos uma informação.")
    for chave, mapa in usados.items():
        if mapa.coluna:
            try:
                indice_coluna(mapa.coluna)
            except ErroDeValidacao as falha:
                raise ErroDeValidacao(f"{POR_CHAVE[chave].rotulo}: {falha}") from falha
    identificam = ("cliente_codigo", "cliente_documento", "cliente_nome")
    if not any(usados.get(chave) and usados[chave].coluna for chave in identificam):
        raise ErroDeValidacao(
            "Informe a coluna do código, do CPF/CNPJ ou do nome do cliente: é por ela "
            "que o sistema sabe de quem é cada linha."
        )
    if layout.linha_inicial < 1:
        raise ErroDeValidacao("A linha onde começam os dados deve ser 1 ou mais.")


def listar_layouts(sessao: Session) -> dict[str, Layout]:
    bruto = configuracoes.obter(sessao, CHAVE_CONFIGURACAO, "")
    try:
        dados = json.loads(bruto) if bruto else {}
    except json.JSONDecodeError:
        return {}
    if not isinstance(dados, dict):
        return {}
    layouts = {}
    for nome, item in dados.items():
        if isinstance(item, dict):
            layout = Layout.de_dict({**item, "nome": nome})
            layouts[layout.nome] = layout
    return dict(sorted(layouts.items(), key=lambda par: normalizar_busca(par[0])))


def salvar_layout(sessao: Session, layout: Layout) -> None:
    nome = (layout.nome or "").strip()
    if not nome:
        raise ErroDeValidacao("Dê um nome ao layout para salvar.", "nome")
    validar_layout(layout)
    layouts = listar_layouts(sessao)
    layouts[nome] = replace(layout, nome=nome)
    _gravar_layouts(sessao, layouts)


def excluir_layout(sessao: Session, nome: str) -> None:
    layouts = listar_layouts(sessao)
    if layouts.pop(nome, None) is None:
        raise ErroDeValidacao(f"O layout {nome!r} não existe.")
    _gravar_layouts(sessao, layouts)


def _gravar_layouts(sessao: Session, layouts: dict[str, Layout]) -> None:
    configuracoes.gravar(
        sessao,
        CHAVE_CONFIGURACAO,
        json.dumps(
            {nome: layout.para_dict() for nome, layout in layouts.items()},
            ensure_ascii=False,
        ),
    )


# ------------------------------------------------------------- planilha-modelo


def exportar_modelo(layout: Layout, destino: Path | str) -> Path:
    """Planilha para preencher: o cabeçalho em cada coluna do layout e as instruções.

    O layout vai junto numa aba oculta, para poder ser aberto de novo no sistema.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    validar_layout(layout)
    arquivo = Path(destino)
    livro = Workbook()
    dados = livro.active
    dados.title = (layout.aba or "Dados")[:31]

    por_coluna: dict[int, list[Campo]] = {}
    for chave, mapa in layout.usados().items():
        if mapa.coluna:
            por_coluna.setdefault(indice_coluna(mapa.coluna), []).append(POR_CHAVE[chave])

    linha_cabecalho = layout.linha_inicial - 1
    if linha_cabecalho >= 1:
        for coluna, campos in sorted(por_coluna.items()):
            celula = dados.cell(
                row=linha_cabecalho, column=coluna,
                value=" / ".join(_titulo_coluna(campo) for campo in campos),
            )
            celula.font = Font(bold=True)
            celula.fill = PatternFill("solid", fgColor="EEF1F5")
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            dados.column_dimensions[get_column_letter(coluna)].width = max(
                14, min(40, len(celula.value) + 4)
            )
        dados.freeze_panes = dados.cell(row=layout.linha_inicial, column=1)

    instrucoes = livro.create_sheet("Instruções")
    instrucoes.append([f"Layout: {layout.nome or '(sem nome)'}"])
    instrucoes["A1"].font = Font(bold=True, size=13)
    instrucoes.append([
        f"Preencha a aba {dados.title!r} a partir da linha {layout.linha_inicial}, "
        "uma linha por honorário (ou por cliente, se for só cadastro). O mesmo "
        "cliente pode aparecer em várias linhas: ele não é cadastrado duas vezes."
    ])
    instrucoes.append([])
    instrucoes.append(["Grupo", "Informação", "Coluna", "Valor fixo (se vazio)", "O que aceita"])
    for celula in instrucoes[4]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="EEF1F5")
    for campo in CAMPOS:
        mapa = layout.campos.get(campo.chave)
        if mapa is None or mapa.vazio:
            continue
        instrucoes.append(
            [campo.grupo, campo.rotulo, mapa.coluna or "—", mapa.padrao, campo.dica]
        )
    for coluna, largura in zip("ABCDE", (30, 30, 9, 24, 60), strict=True):
        instrucoes.column_dimensions[coluna].width = largura

    oculta = livro.create_sheet(ABA_LAYOUT)
    oculta["A1"] = json.dumps(layout.para_dict(), ensure_ascii=False)
    oculta.sheet_state = "hidden"

    livro.save(arquivo)
    return arquivo


def _titulo_coluna(campo: Campo) -> str:
    if campo.grupo in (GRUPO_CLIENTE, GRUPO_VINCULO):
        return campo.rotulo
    prefixo = {GRUPO_CONTATO: "Contato", GRUPO_CONTRATO: "Contrato",
               GRUPO_HONORARIO: "Honorário"}[campo.grupo]
    rotulo = campo.rotulo.removesuffix(" do contato").removesuffix(" do contrato")
    rotulo = rotulo.removesuffix(" do honorário")
    return f"{prefixo}: {rotulo}"


def ler_layout_do_modelo(arquivo: Path | str) -> Layout:
    """Abre o layout que vai na aba oculta da planilha-modelo."""
    from openpyxl import load_workbook

    livro = load_workbook(Path(arquivo), read_only=True)
    try:
        if ABA_LAYOUT not in livro.sheetnames:
            raise ErroDeValidacao(
                "Esta planilha não foi gerada pelo botão Exportar modelo: não tem o layout."
            )
        bruto = next(livro[ABA_LAYOUT].iter_rows(max_row=1, values_only=True), (None,))[0]
    finally:
        livro.close()
    try:
        return Layout.de_dict(json.loads(bruto or "{}"))
    except (json.JSONDecodeError, AttributeError) as falha:
        raise ErroDeValidacao("O layout guardado na planilha está corrompido.") from falha


# ----------------------------------------------------------- leitura de valores

_SIM = {"sim", "s", "x", "1", "true", "verdadeiro", "yes"}
_NAO = {"nao", "n", "0", "false", "falso", "no"}


def _apelidos(enumeracao, extras: dict[str, object] | None = None) -> dict[str, object]:
    mapa = {}
    for item in enumeracao:
        mapa[normalizar_busca(item.value)] = item
        mapa[normalizar_busca(item.value.replace("_", " "))] = item
        mapa[normalizar_busca(item.rotulo)] = item
    mapa.update(extras or {})
    return mapa


APELIDOS = {
    "pessoa": _apelidos(TipoPessoa, {
        "fisica": TipoPessoa.PF, "juridica": TipoPessoa.PJ, "cpf": TipoPessoa.PF,
        "cnpj": TipoPessoa.PJ, "empresa": TipoPessoa.PJ, "domestica": TipoPessoa.PF,
    }),
    "subtipo": _apelidos(SubtipoCliente, {"": SubtipoCliente.NENHUM, "-": SubtipoCliente.NENHUM,
                                          "nenhum": SubtipoCliente.NENHUM}),
    "situacao": _apelidos(SituacaoCliente),
    "periodicidade": _apelidos(Periodicidade),
    "meio": _apelidos(MeioCobranca),
    "envio": _apelidos(SituacaoEnvio, {
        "enviado": SituacaoEnvio.ENVIADO, "sim": SituacaoEnvio.ENVIADO,
        "nao": SituacaoEnvio.NAO_ENVIADO,
    }),
    "status": _apelidos(StatusPagamento, {
        chave: valores[0] for chave, valores in MAPA_STATUS.items()
    } | {"aberto": StatusPagamento.EM_ABERTO, "quitado": StatusPagamento.PAGO}),
}


@dataclass
class Referencias:
    """Cadastros que a leitura consulta: tipos de honorário e formas de envio."""

    tipos: dict[str, TipoHonorario]
    formas: dict[str, FormaEnvio]

    @classmethod
    def do_banco(cls, sessao: Session) -> Referencias:
        return cls(
            tipos={
                normalizar_busca(tipo.nome): tipo
                for tipo in sessao.scalars(select(TipoHonorario)).unique().all()
            },
            formas={
                normalizar_busca(forma.nome): forma
                for forma in sessao.scalars(select(FormaEnvio)).all()
            },
        )


class _ValorInvalido(Exception):
    pass


def converter(campo: Campo, bruto, referencias: Referencias):
    """Valor da célula (ou o fixo do layout) para o tipo do campo.

    Devolve ``None`` para vazio e levanta ``_ValorInvalido`` com o motivo.
    """
    if bruto is None or (isinstance(bruto, str) and not bruto.strip()):
        return None
    tipo = campo.tipo
    if tipo == "texto":
        texto = _texto(bruto)
        if isinstance(bruto, float) and bruto.is_integer():
            texto = str(int(bruto))   # código 302 lido como 302.0
        return texto or None
    if tipo == "documento":
        digitos = so_digitos(_texto(int(bruto) if isinstance(bruto, float) else bruto))
        if len(digitos) in (10, 13):   # o Excel come o zero à esquerda
            digitos = "0" + digitos
        if not validar_documento(digitos):
            raise _ValorInvalido("CPF/CNPJ inválido")
        return digitos
    if tipo == "telefone":
        telefone = normalizar_telefone(_texto(int(bruto) if isinstance(bruto, float) else bruto))
        return telefone or None
    if tipo == "dinheiro":
        valor = _decimal(bruto)
        if valor is None:
            raise _ValorInvalido("Valor em dinheiro não reconhecido")
        return valor
    if tipo == "taxa":
        valor = _juros(bruto)
        if valor is None or valor < 0:
            raise _ValorInvalido("Juros por dia não reconhecido")
        return valor
    if tipo == "data":
        valor = _data(bruto)
        if valor is None:
            raise _ValorInvalido("Data não reconhecida (use dd/mm/aaaa)")
        return valor
    if tipo == "inteiro":
        try:
            return int(Decimal(str(bruto).strip().replace(",", ".")))
        except Exception as falha:  # noqa: BLE001
            raise _ValorInvalido("Número inteiro não reconhecido") from falha
    if tipo == "competencia":
        try:
            competencia = Competencia.de_texto(bruto)
        except ValueError as falha:
            raise _ValorInvalido(str(falha)) from falha
        if competencia.vazia:
            raise _ValorInvalido("Competência não reconhecida (use mm/aaaa ou o ano)")
        return competencia
    if tipo == "sim_nao":
        alvo = normalizar_busca(_texto(bruto))
        if alvo in _SIM:
            return True
        if alvo in _NAO:
            return False
        raise _ValorInvalido("Use Sim ou Não")
    if tipo == "tipo_honorario":
        tipo_honorario = referencias.tipos.get(normalizar_busca(_texto(bruto)))
        if tipo_honorario is None:
            raise _ValorInvalido("Tipo de honorário não existe no cadastro")
        return tipo_honorario
    if tipo == "forma_envio":
        forma = referencias.formas.get(normalizar_busca(_texto(bruto)))
        if forma is None:
            raise _ValorInvalido("Forma de envio não existe no cadastro")
        return forma
    alvo = normalizar_busca(_texto(bruto))
    item = APELIDOS[tipo].get(alvo)
    if item is None:
        raise _ValorInvalido("Valor não reconhecido")
    return item


# ------------------------------------------------------------------- a prévia


@dataclass
class LinhaPersonalizada:
    linha: int
    valores: dict[str, object] = field(default_factory=dict)
    # o que a leitura achou (fixo) e o que a conferência com o banco achou
    # (refeito a cada conferência, por isso separado)
    leitura: list[Problema] = field(default_factory=list)
    conferencia: list[Problema] = field(default_factory=list)
    chave_cliente: str = ""
    acao_cliente: str = ""
    acao_contrato: str = ""
    acao_honorario: str = ""
    importar: bool = True

    @property
    def problemas(self) -> list[Problema]:
        return [*self.leitura, *self.conferencia]

    @property
    def bloqueada(self) -> bool:
        return any(problema.bloqueia for problema in self.problemas)

    @property
    def rotulo_cliente(self) -> str:
        codigo = self.valores.get("cliente_codigo") or ""
        nome = self.valores.get("cliente_nome") or self.valores.get("cliente_documento") or ""
        return f"{codigo} – {nome}" if codigo and nome else str(codigo or nome)

    def tem(self, *chaves: str) -> bool:
        return any(self.valores.get(chave) is not None for chave in chaves)


@dataclass
class ClientePlanejado:
    chave: str
    valores: dict[str, object] = field(default_factory=dict)
    contatos: list[dict[str, str]] = field(default_factory=list)
    linhas: list[int] = field(default_factory=list)
    existente_id: int | None = None
    # "documento", "codigo" ou "nome": só o CPF/CNPJ autoriza trocar o nome
    reconhecido_por: str = ""


@dataclass
class Previa:
    arquivo: Path
    layout: Layout
    aba: str = ""
    linhas: list[LinhaPersonalizada] = field(default_factory=list)
    problemas: list[Problema] = field(default_factory=list)

    @property
    def importaveis(self) -> list[LinhaPersonalizada]:
        return [linha for linha in self.linhas if linha.importar and not linha.bloqueada]

    def todos_problemas(self) -> list[Problema]:
        return [*self.problemas, *(p for linha in self.linhas for p in linha.problemas)]

    def resumo(self) -> str:
        validas = [linha for linha in self.linhas if not linha.bloqueada]
        novos = len({
            linha.chave_cliente for linha in validas if linha.acao_cliente.startswith("Novo")
        })
        atualizar = len({
            linha.chave_cliente for linha in validas
            if linha.acao_cliente.startswith("Atualizar")
        })
        honorarios = [linha for linha in validas if linha.acao_honorario.startswith("Novo")]
        total = sum((linha.valores["honorario_valor"] for linha in honorarios), start=ZERO)
        contratos = sum(1 for linha in validas if linha.acao_contrato.startswith(("Novo", "Atual",
                                                                                   "Reaj")))
        erros = sum(1 for linha in self.linhas if linha.bloqueada)
        return (
            f"{len(self.linhas)} linha(s): {novos} cliente(s) novo(s), {atualizar} a "
            f"atualizar, {contratos} contrato(s) e {len(honorarios)} honorário(s) "
            f"({formatar_moeda(total)}). {erros} linha(s) com erro."
        )


def ler(sessao: Session, arquivo: Path | str, layout: Layout,
        atualizar_existentes: bool = True) -> Previa:
    """Lê a planilha com o layout e confere tudo contra o banco. Não grava nada."""
    from openpyxl import load_workbook

    validar_layout(layout)
    caminho = Path(arquivo)
    previa = Previa(arquivo=caminho, layout=layout)
    referencias = Referencias.do_banco(sessao)

    livro = load_workbook(caminho, data_only=True, read_only=True)
    try:
        if layout.aba:
            nomes = {normalizar_busca(nome): nome for nome in livro.sheetnames}
            nome_aba = nomes.get(normalizar_busca(layout.aba))
            if nome_aba is None:
                raise ErroDeValidacao(
                    f"A planilha não tem a aba {layout.aba!r}. Abas: "
                    + ", ".join(n for n in livro.sheetnames if n != ABA_LAYOUT)
                )
        else:
            nome_aba = next(n for n in livro.sheetnames if n != ABA_LAYOUT)
        previa.aba = nome_aba
        colunas = {
            chave: indice_coluna(mapa.coluna)
            for chave, mapa in layout.usados().items() if mapa.coluna
        }
        maior = max(colunas.values())
        for numero, celulas in enumerate(
            livro[nome_aba].iter_rows(min_row=layout.linha_inicial, max_col=maior,
                                      values_only=True),
            start=layout.linha_inicial,
        ):
            brutos = {
                chave: celulas[indice - 1] if indice - 1 < len(celulas) else None
                for chave, indice in colunas.items()
            }
            if all(valor in (None, "") or (isinstance(valor, str) and not valor.strip())
                   for valor in brutos.values()):
                continue
            previa.linhas.append(_ler_linha(previa, numero, brutos, referencias))
    finally:
        livro.close()

    if not previa.linhas:
        previa.problemas.append(Problema(
            previa.aba, 0, "arquivo", caminho.name,
            f"Nenhuma linha com dados a partir da linha {layout.linha_inicial}", "ERRO",
        ))
    conferir(sessao, previa, atualizar_existentes)
    return previa


def _ler_linha(previa: Previa, numero: int, brutos: dict, referencias: Referencias
               ) -> LinhaPersonalizada:
    linha = LinhaPersonalizada(linha=numero)
    aba = previa.aba
    for chave, mapa in previa.layout.usados().items():
        campo = POR_CHAVE[chave]
        bruto = brutos.get(chave)
        origem = "planilha"
        if bruto is None or (isinstance(bruto, str) and not bruto.strip()):
            bruto, origem = (mapa.padrao or None), "valor fixo do layout"
        try:
            valor = converter(campo, bruto, referencias)
        except _ValorInvalido as falha:
            # dado errado em campo de identidade ou de dinheiro trava a linha;
            # nos demais o valor é ignorado com aviso
            critico = campo.tipo in ("documento", "dinheiro", "data", "taxa", "competencia",
                                     "tipo_honorario")
            linha.leitura.append(Problema(
                aba, numero, campo.rotulo, _texto(bruto)[:120],
                f"{falha} ({origem})" + ("" if critico else ": ignorado"),
                "ERRO" if critico else "AVISO",
            ))
            continue
        if valor is not None:
            linha.valores[chave] = valor
    _validar_linha(linha, aba)
    return linha


def _validar_linha(linha: LinhaPersonalizada, aba: str) -> None:
    valores = linha.valores
    numero = linha.linha

    def erro(campo: str, motivo: str) -> None:
        linha.leitura.append(Problema(aba, numero, campo, "", motivo, "ERRO"))

    if not linha.tem("cliente_codigo", "cliente_documento", "cliente_nome"):
        erro("Cliente", "Linha sem código, CPF/CNPJ ou nome do cliente")
        return

    documento = valores.get("cliente_documento")
    pelo_documento = None
    if documento:
        pelo_documento = TipoPessoa.PF if len(documento) == 11 else TipoPessoa.PJ
    tipo = valores.get("cliente_tipo_pessoa")
    if pelo_documento and tipo and tipo is not pelo_documento:
        linha.leitura.append(Problema(
            aba, numero, "Tipo de pessoa", tipo.value,
            f"O CPF/CNPJ é de {pelo_documento.rotulo}: tipo corrigido",
        ))
        valores["cliente_tipo_pessoa"] = pelo_documento
    elif pelo_documento and not tipo:
        valores["cliente_tipo_pessoa"] = pelo_documento

    if linha.tem("contrato_valor"):
        if valores["contrato_valor"] <= 0:
            erro("Valor do contrato", "O valor do contrato deve ser maior que zero")
        tipo_contrato = valores.get("contrato_tipo")
        if tipo_contrato is None:
            erro("Tipo de honorário (contrato)", "Contrato sem tipo de honorário")
        elif (valores.get("contrato_periodicidade") or tipo_contrato.periodicidade) is (
            Periodicidade.AVULSO
        ):
            erro("Periodicidade", f"{tipo_contrato.nome} é avulso: não tem contrato")
        dia = valores.get("contrato_dia_vencimento")
        if dia is not None and not 1 <= dia <= 31:
            erro("Dia do vencimento", "O dia do vencimento deve estar entre 1 e 31")
    elif linha.tem("contrato_tipo") and any(
        linha.tem(campo.chave) for campo in campos_do_grupo(GRUPO_CONTRATO)
        if campo.chave not in ("contrato_tipo", "contrato_valor")
    ):
        linha.leitura.append(Problema(
            aba, numero, "Valor do contrato", "",
            "Dados de contrato sem o valor: o contrato só é criado ou mudado com o valor",
        ))

    if linha.tem("honorario_valor"):
        if valores["honorario_valor"] <= 0:
            erro("Valor do honorário", "O valor do honorário deve ser maior que zero")
        tipo_honorario = valores.get("honorario_tipo")
        if tipo_honorario is None:
            erro("Tipo de honorário", "Honorário sem tipo de honorário")
        if not linha.tem("honorario_vencimento"):
            erro("Vencimento", "Honorário sem vencimento: não é possível calcular atraso")
        if tipo_honorario is not None:
            competencia = valores.get("honorario_competencia") or Competencia()
            exige = tipo_honorario.exige_competencia
            if exige is ExigeCompetencia.MES and not (competencia.ano and competencia.mes):
                erro("Competência", f"{tipo_honorario.nome} exige competência mm/aaaa")
            elif exige is ExigeCompetencia.ANO and not competencia.ano:
                erro("Competência", f"{tipo_honorario.nome} exige o ano de competência")
            elif exige is ExigeCompetencia.NENHUMA and not valores.get("honorario_descricao"):
                erro("Descrição", f"{tipo_honorario.nome} não usa competência: descreva o serviço")
            if exige is ExigeCompetencia.ANO and competencia.mes:
                valores["honorario_competencia"] = Competencia(ano=competencia.ano)
            elif exige is ExigeCompetencia.NENHUMA and not competencia.vazia:
                valores.pop("honorario_competencia", None)
    elif linha.tem("honorario_vencimento", "honorario_competencia"):
        linha.leitura.append(Problema(
            aba, numero, "Valor do honorário", "",
            "Dados de honorário sem o valor: honorário não lançado",
        ))


# ---------------------------------------------------------------- conferência


CAMPOS_CLIENTE = (
    "cliente_codigo", "cliente_nome", "cliente_documento", "cliente_tipo_pessoa",
    "cliente_subtipo", "cliente_situacao", "cliente_juros_dia", "cliente_observacoes",
    "empresa_vinculo", "suspensao_motivo", "suspensao_ate",
)


def _agrupar_clientes(aba: str, linhas: list[LinhaPersonalizada], avisar: bool = True
                      ) -> dict[str, ClientePlanejado]:
    """Junta as linhas do mesmo cliente: pelo documento, pelo código ou pelo nome.

    Com ``avisar``, as divergências entre linhas do mesmo cliente vão para a
    conferência de cada linha.
    """
    planos: dict[str, ClientePlanejado] = {}
    por_documento: dict[str, str] = {}
    por_codigo: dict[str, str] = {}
    por_nome: dict[str, str] = {}
    for linha in linhas:
        valores = linha.valores
        documento = valores.get("cliente_documento") or ""
        codigo = valores.get("cliente_codigo") or ""
        nome = normalizar_busca(valores.get("cliente_nome") or "")
        chave = (
            (documento and por_documento.get(documento))
            or (codigo and por_codigo.get(codigo))
            or ""
        )
        if not chave and nome in por_nome:
            # pelo nome só quando documento e código não desmentem
            candidato = planos[por_nome[nome]].valores
            if (not documento or candidato.get("cliente_documento") in (None, documento)) and (
                not codigo or candidato.get("cliente_codigo") in (None, codigo)
            ):
                chave = por_nome[nome]
        chave = chave or (f"doc:{documento}" if documento else "") or (
            f"cod:{codigo}" if codigo else f"nome:{nome}"
        )
        plano = planos.get(chave)
        if plano is None:
            plano = planos[chave] = ClientePlanejado(chave=chave)
        else:
            outro_documento = plano.valores.get("cliente_documento")
            if documento and outro_documento and documento != outro_documento:
                # mesmo código com documentos diferentes: são clientes diferentes
                if avisar:
                    linha.conferencia.append(Problema(
                        aba, linha.linha, "CPF/CNPJ", documento,
                        f"O código {codigo} aparece antes com outro CPF/CNPJ na planilha",
                        "ERRO",
                    ))
                continue
        for chave_campo in CAMPOS_CLIENTE:
            valor = valores.get(chave_campo)
            if valor is None:
                continue
            atual = plano.valores.get(chave_campo)
            if atual is None:
                plano.valores[chave_campo] = valor
            elif atual != valor and chave_campo != "cliente_observacoes":
                if not avisar:
                    continue
                campo = POR_CHAVE[chave_campo]
                linha.conferencia.append(Problema(
                    aba, linha.linha, campo.rotulo, _texto(valor)[:120],
                    f"Difere da linha {plano.linhas[0]} para o mesmo cliente: mantido "
                    f"{_texto(_mostrar(atual))[:60]}",
                ))
            elif chave_campo == "cliente_observacoes" and valor not in atual:
                plano.valores[chave_campo] = f"{atual}\n{valor}"
        contato = {
            "nome": valores.get("contato_nome") or "",
            "nome_tratamento": valores.get("contato_tratamento") or "",
            "telefone": valores.get("contato_telefone") or "",
            "email": valores.get("contato_email") or "",
        }
        if any(contato.values()) and not any(
            _mesmo_contato(contato, outro) for outro in plano.contatos
        ):
            plano.contatos.append(contato)
        plano.linhas.append(linha.linha)
        linha.chave_cliente = chave
        if documento:
            por_documento.setdefault(documento, chave)
        if codigo:
            por_codigo.setdefault(codigo, chave)
        if nome:
            por_nome.setdefault(nome, chave)
    return planos


def _mostrar(valor) -> str:
    if hasattr(valor, "rotulo") and not isinstance(valor, Competencia):
        return valor.rotulo
    if isinstance(valor, (TipoHonorario, FormaEnvio)):
        return valor.nome
    if isinstance(valor, date):
        return f"{valor:%d/%m/%Y}"
    return str(valor)


def _mesmo_contato(novo: dict[str, str], outro: dict[str, str]) -> bool:
    if novo["telefone"] and novo["telefone"] == outro["telefone"]:
        return True
    if novo["email"] and novo["email"].lower() == (outro["email"] or "").lower():
        return True
    if not novo["telefone"] and not novo["email"]:
        return bool(novo["nome"]) and normalizar_busca(novo["nome"]) == normalizar_busca(
            outro["nome"] or ""
        )
    return False


def _cliente_no_banco(sessao: Session, plano: ClientePlanejado
                      ) -> tuple[Cliente | None, str]:
    """Acha o cliente já cadastrado. Devolve (cliente, conflito)."""
    valores = plano.valores
    documento = valores.get("cliente_documento")
    codigo = valores.get("cliente_codigo")
    plano.reconhecido_por = ""
    if documento:
        achado = servico_clientes.por_documento(sessao, documento)
        if achado is not None:
            plano.reconhecido_por = "documento"
            return achado, ""
    if codigo:
        achado = servico_clientes.por_codigo(sessao, codigo)
        if achado is not None:
            if documento and achado.documento and achado.documento != documento:
                return None, (
                    f"O código {codigo} já é de {achado.nome}, com outro CPF/CNPJ"
                )
            plano.reconhecido_por = "codigo"
            return achado, ""
    nome = normalizar_busca(valores.get("cliente_nome") or "")
    if nome:
        achados = list(sessao.scalars(
            select(Cliente).where(Cliente.nome_busca == nome, Cliente.excluido_em.is_(None))
        ).unique())
        # pelo nome só quando é um cliente só e o CPF/CNPJ não desmente
        compativeis = [
            c for c in achados if not documento or not c.documento or c.documento == documento
        ]
        if len(compativeis) == 1:
            plano.reconhecido_por = "nome"
            return compativeis[0], ""
    return None, ""


def _dados_de(cliente: Cliente) -> servico_clientes.DadosCliente:
    """O cadastro atual do cliente no formato que o serviço de clientes grava."""
    return servico_clientes.DadosCliente(
        codigo=cliente.codigo,
        nome=cliente.nome,
        documento=cliente.documento,
        tipo_pessoa=cliente.tipo_pessoa,
        subtipo=cliente.subtipo,
        situacao=cliente.situacao,
        observacoes=cliente.observacoes,
        juros_dia=servico_clientes.juros_atual(cliente) if cliente.juros else None,
        contatos=[
            servico_clientes.DadosContato(
                id=c.id, nome=c.nome, nome_tratamento=c.nome_tratamento,
                telefone=c.telefone, email=c.email, principal=c.principal,
                ativo=c.ativo, observacoes=c.observacoes,
            )
            for c in cliente.contatos
        ],
        empresa_id=cliente.empresa_id,
    )


def _mesclar(atual: servico_clientes.DadosCliente, plano: ClientePlanejado,
             empresa_id: int | None) -> tuple[servico_clientes.DadosCliente, list[str]]:
    """Aplica ao cadastro só o que a planilha traz de novo ou de diferente."""
    valores = plano.valores
    novo = replace(atual, contatos=[replace(c) for c in atual.contatos or []])
    mudou: list[str] = []
    # só acento ou maiúscula diferente não é nome novo; e só se troca o nome de
    # cliente reconhecido pelo CPF/CNPJ (o código pode ser de outro sistema)
    nome = (valores.get("cliente_nome") or "").strip()
    if nome and normalizar_busca(nome) != normalizar_busca(atual.nome) and (
        not atual.nome or plano.reconhecido_por == "documento"
    ):
        novo.nome = nome
        mudou.append("nome")
    if valores.get("cliente_documento") and not atual.documento:
        novo.documento = valores["cliente_documento"]
        mudou.append("CPF/CNPJ")
    for chave, atributo, rotulo in (
        ("cliente_tipo_pessoa", "tipo_pessoa", "tipo de pessoa"),
        ("cliente_subtipo", "subtipo", "subtipo"),
        ("cliente_situacao", "situacao", "situação"),
        ("cliente_juros_dia", "juros_dia", "juros/dia"),
    ):
        valor = valores.get(chave)
        if valor is not None and valor != getattr(atual, atributo):
            setattr(novo, atributo, valor)
            mudou.append(rotulo)
    if novo.tipo_pessoa is TipoPessoa.PJ and novo.subtipo is not SubtipoCliente.NENHUM:
        novo.subtipo = SubtipoCliente.NENHUM
    observacao = valores.get("cliente_observacoes")
    if observacao and observacao not in (atual.observacoes or ""):
        novo.observacoes = f"{atual.observacoes}\n{observacao}".strip()
        mudou.append("observações")
    if empresa_id is not None and empresa_id != atual.empresa_id:
        novo.empresa_id = empresa_id
        mudou.append("empresa que paga")

    for contato in plano.contatos:
        existentes = [
            {"nome": c.nome, "nome_tratamento": c.nome_tratamento,
             "telefone": c.telefone, "email": c.email}
            for c in novo.contatos
        ]
        indice = next(
            (i for i, outro in enumerate(existentes) if _mesmo_contato(contato, outro)), None
        )
        if indice is None:
            # o cadastro descarta contato sem nome e sem telefone (só com e-mail)
            nome_contato = contato["nome"] or contato["nome_tratamento"] or (
                "" if contato["telefone"] else novo.nome
            )
            novo.contatos.append(servico_clientes.DadosContato(
                nome=nome_contato, nome_tratamento=contato["nome_tratamento"],
                telefone=contato["telefone"], email=contato["email"],
                principal=not novo.contatos,
            ))
            mudou.append("novo contato")
            continue
        alvo = novo.contatos[indice]
        completou = False
        for atributo in ("nome", "nome_tratamento", "telefone", "email"):
            if contato[atributo] and not getattr(alvo, atributo):
                setattr(alvo, atributo, contato[atributo])
                completou = True
        if completou:
            mudou.append("contato")
    return novo, list(dict.fromkeys(mudou))


def _resolver_empresa(sessao: Session, referencia: str, planos: dict[str, ClientePlanejado]
                      ) -> tuple[Cliente | None, str | None]:
    """A empresa do vínculo: no banco (código ou CNPJ) ou entre os clientes da planilha.

    Devolve (cliente do banco, chave do plano da planilha).
    """
    digitos = so_digitos(referencia)
    if len(digitos) == 13:
        digitos = "0" + digitos
    achado = None
    if len(digitos) == 14:
        achado = servico_clientes.por_documento(sessao, digitos)
    if achado is None:
        achado = servico_clientes.por_codigo(sessao, referencia)
    if achado is not None:
        return achado, None
    for chave, plano in planos.items():
        if (len(digitos) == 14 and plano.valores.get("cliente_documento") == digitos) or (
            plano.valores.get("cliente_codigo") == referencia.strip()
        ):
            return None, chave
    return None, None


def conferir(sessao: Session, previa: Previa, atualizar_existentes: bool = True
             ) -> dict[str, ClientePlanejado]:
    """Decide o que cada linha vai fazer (cliente novo, atualizar, honorário repetido…).

    Pode ser chamada de novo a qualquer momento: refaz a conferência do zero.
    """
    for linha in previa.linhas:
        linha.conferencia.clear()
        linha.chave_cliente = ""
        linha.acao_cliente = linha.acao_contrato = linha.acao_honorario = ""
    validas = [linha for linha in previa.linhas if not linha.bloqueada]
    planos = _agrupar_clientes(previa.aba, validas)
    por_chave: dict[str, list[LinhaPersonalizada]] = {}
    for linha in validas:
        if linha.chave_cliente:
            por_chave.setdefault(linha.chave_cliente, []).append(linha)

    for chave, plano in planos.items():
        linhas = por_chave.get(chave, [])
        if not linhas:
            continue

        def avisar(campo: str, valor: str, motivo: str, gravidade: str = "AVISO",
                   todas: bool = False, _linhas=linhas) -> None:
            for linha in (_linhas if todas else _linhas[:1]):
                linha.conferencia.append(
                    Problema(previa.aba, linha.linha, campo, valor, motivo, gravidade)
                )

        existente, conflito = _cliente_no_banco(sessao, plano)
        if conflito:
            avisar("Cliente", plano.valores.get("cliente_codigo", ""), conflito, "ERRO", True)
            continue
        if existente is not None:
            plano.existente_id = existente.id
            codigo = plano.valores.get("cliente_codigo")
            if codigo and codigo != existente.codigo:
                avisar("Código", codigo,
                       f"Cliente já cadastrado com o código {existente.codigo}: mantido")
            nome = plano.valores.get("cliente_nome") or ""
            if (plano.reconhecido_por == "codigo" and nome
                    and normalizar_busca(nome) != existente.nome_busca):
                avisar("Nome / razão social", nome,
                       f"O código {existente.codigo} é de {existente.nome} no cadastro: "
                       "nome mantido. Confira se é o mesmo cliente antes de importar")
            tipo = plano.valores.get("cliente_tipo_pessoa") or existente.tipo_pessoa
            empresa_id = _empresa_planejada(sessao, plano, planos, tipo, avisar)
            _, mudou = _mesclar(_dados_de(existente), plano, empresa_id)
            if plano.valores.get("suspensao_motivo") and not existente.cobranca_suspensa:
                mudou.append("suspensão da cobrança")
            if not mudou:
                acao = f"Já cadastrado ({existente.codigo})"
            elif atualizar_existentes:
                acao = f"Atualizar {existente.codigo}: " + ", ".join(mudou)
            else:
                acao = f"Já cadastrado ({existente.codigo}), sem atualizar"
        else:
            if not plano.valores.get("cliente_nome"):
                avisar("Cliente", plano.valores.get("cliente_codigo")
                       or plano.valores.get("cliente_documento", ""),
                       "Cliente não está cadastrado e a linha não traz o nome", "ERRO", True)
                continue
            tipo = plano.valores.get("cliente_tipo_pessoa") or TipoPessoa.PJ
            _empresa_planejada(sessao, plano, planos, tipo, avisar)
            acao = "Novo cliente"
            if plano.valores.get("cliente_juros_dia") is None:
                avisar("Juros por dia (R$)", "",
                       "Cliente novo sem juros por dia: cadastrado com R$ 0,00")
        for linha in linhas:
            linha.acao_cliente = acao

    _conferir_contratos(sessao, previa, planos, atualizar_existentes)
    _conferir_honorarios(sessao, previa, planos)
    return planos


def _empresa_planejada(sessao: Session, plano: ClientePlanejado,
                       planos: dict[str, ClientePlanejado], tipo: TipoPessoa, avisar
                       ) -> int | None:
    """Id da empresa do vínculo; -1 quando ela é nova, da própria planilha."""
    referencia = plano.valores.get("empresa_vinculo")
    if not referencia:
        return None
    if tipo is TipoPessoa.PJ:
        avisar("Empresa que paga por ele", referencia,
               "Só pessoa física é vinculada a uma empresa: vínculo ignorado")
        return None
    empresa, chave = _resolver_empresa(sessao, referencia, planos)
    if empresa is None and chave is None:
        avisar("Empresa que paga por ele", referencia,
               "Empresa não encontrada no cadastro nem na planilha: vínculo ignorado")
        return None
    # empresa nova (da planilha) ainda não tem id: o vínculo é feito depois de criá-la
    return empresa.id if empresa is not None else -1


def _conferir_contratos(sessao: Session, previa: Previa, planos: dict[str, ClientePlanejado],
                        atualizar_existentes: bool = True) -> None:
    vistos: dict[tuple[str, int], LinhaPersonalizada] = {}
    for linha in previa.linhas:
        if linha.bloqueada or not linha.tem("contrato_valor"):
            continue
        tipo = linha.valores["contrato_tipo"]
        chave = (linha.chave_cliente, tipo.id)
        if chave in vistos:
            linha.acao_contrato = f"Já na linha {vistos[chave].linha}"
            continue
        vistos[chave] = linha
        plano = planos[linha.chave_cliente]
        existente = _contrato_existente(sessao, plano.existente_id, tipo.id)
        if existente is None:
            valor = formatar_moeda(linha.valores["contrato_valor"])
            linha.acao_contrato = f"Novo: {tipo.nome} {valor}"
            continue
        mudancas = _mudancas_contrato(existente, linha.valores)
        if not atualizar_existentes:
            linha.acao_contrato = "Já cadastrado, sem atualizar"
        elif linha.valores["contrato_valor"] != existente.valor:
            linha.acao_contrato = (
                f"Reajuste {formatar_moeda(existente.valor)} → "
                f"{formatar_moeda(linha.valores['contrato_valor'])}"
            )
            if mudancas:
                linha.acao_contrato += " e " + ", ".join(mudancas)
        elif mudancas:
            linha.acao_contrato = "Atualizar: " + ", ".join(mudancas)
        else:
            linha.acao_contrato = "Já cadastrado"


def _contrato_existente(sessao: Session, cliente_id: int | None, tipo_id: int
                        ) -> Contrato | None:
    if cliente_id is None:
        return None
    return sessao.scalar(select(Contrato).where(
        Contrato.cliente_id == cliente_id,
        Contrato.tipo_id == tipo_id,
        Contrato.excluido_em.is_(None),
    ))


_CAMPOS_CONTRATO = (
    ("contrato_periodicidade", "periodicidade", "periodicidade"),
    ("contrato_dia_vencimento", "dia_vencimento", "dia do vencimento"),
    ("contrato_meses_vencimento", "meses_vencimento", "meses até vencer"),
    ("contrato_meses_referencia", "meses_referencia", "meses de referência"),
    ("contrato_meio", "meio", "boleto/pix"),
    ("contrato_vigencia_inicio", "vigencia_inicio", "início da vigência"),
    ("contrato_vigencia_fim", "vigencia_fim", "fim da vigência"),
    ("contrato_suspenso", "suspenso", "suspensão"),
    ("contrato_observacoes", "observacoes", "observações"),
)


def _mudancas_contrato(contrato: Contrato, valores: dict) -> list[str]:
    mudou = [
        rotulo for chave, atributo, rotulo in _CAMPOS_CONTRATO
        if valores.get(chave) is not None and valores[chave] != getattr(contrato, atributo)
    ]
    forma = valores.get("contrato_forma_envio")
    if forma is not None and forma.id != contrato.forma_envio_id:
        mudou.append("forma de envio")
    return mudou


def _conferir_honorarios(sessao: Session, previa: Previa, planos: dict[str, ClientePlanejado]
                         ) -> None:
    vistos: dict[tuple, LinhaPersonalizada] = {}
    for linha in previa.linhas:
        if linha.bloqueada or not linha.tem("honorario_valor"):
            continue
        valores = linha.valores
        tipo = valores["honorario_tipo"]
        competencia = valores.get("honorario_competencia") or Competencia()
        chave_competencia = servico_honorarios.chave_competencia(competencia)
        chave = (
            (linha.chave_cliente, tipo.id, chave_competencia) if chave_competencia
            else (linha.chave_cliente, tipo.id, valores["honorario_vencimento"],
                  valores["honorario_valor"])
        )
        if chave in vistos:
            linha.acao_honorario = f"Repetido (linha {vistos[chave].linha}): ignorado"
            continue
        vistos[chave] = linha
        cliente_id = planos[linha.chave_cliente].existente_id
        existente = _honorario_existente(
            sessao, cliente_id, tipo.id, competencia, valores["honorario_vencimento"],
            valores["honorario_valor"], valores.get("honorario_status"),
        )
        if existente is not None:
            linha.acao_honorario = (
                f"Já cadastrado (venc. {existente.vencimento:%d/%m/%Y}): ignorado"
            )
            continue
        status = valores.get("honorario_status") or StatusPagamento.EM_ABERTO
        rotulo = competencia.rotulo() or valores.get("honorario_descricao") or ""
        linha.acao_honorario = (
            f"Novo: {tipo.nome} {rotulo} {formatar_moeda(valores['honorario_valor'])}".replace(
                "  ", " ")
            + ("" if status is StatusPagamento.EM_ABERTO else f" ({status.rotulo.lower()})")
        )


def _honorario_existente(sessao: Session, cliente_id: int | None, tipo_id: int,
                         competencia: Competencia, vencimento: date, valor: Decimal,
                         status: StatusPagamento | None = None) -> Honorario | None:
    """O lançamento que a linha repetiria.

    Cancelado não impede lançar de novo — a não ser que a linha também seja de um
    cancelado: aí é o mesmo, e reimportar a planilha não o duplica.
    """
    if cliente_id is None:
        return None
    consulta = select(Honorario).where(
        Honorario.cliente_id == cliente_id, Honorario.tipo_id == tipo_id
    )
    if not competencia.vazia:
        consulta = consulta.where(
            Honorario.competencia_chave == servico_honorarios.chave_competencia(competencia)
        )
    else:
        consulta = consulta.where(
            Honorario.vencimento == vencimento, Honorario.valor_original == valor
        )
    if status is StatusPagamento.CANCELADO:
        consulta = consulta.where(Honorario.status == StatusPagamento.CANCELADO)
    else:
        consulta = consulta.where(Honorario.status != StatusPagamento.CANCELADO)
    return sessao.scalars(consulta.limit(1)).first()


# -------------------------------------------------------------------- gravação


@dataclass
class Resultado:
    clientes_criados: int = 0
    clientes_atualizados: int = 0
    clientes_sem_alteracao: int = 0
    contratos_criados: int = 0
    contratos_atualizados: int = 0
    honorarios_criados: int = 0
    honorarios_existentes: int = 0
    linhas_ignoradas: int = 0
    falhas: list[Problema] = field(default_factory=list)
    importacao_id: int | None = None

    @property
    def resumo(self) -> str:
        texto = (
            f"{self.clientes_criados} cliente(s) criado(s), {self.clientes_atualizados} "
            f"atualizado(s) e {self.clientes_sem_alteracao} sem alteração; "
            f"{self.contratos_criados} contrato(s) criado(s) e {self.contratos_atualizados} "
            f"atualizado(s); {self.honorarios_criados} honorário(s) lançado(s) "
            f"({self.honorarios_existentes} já existiam). "
            f"{self.linhas_ignoradas} linha(s) ignorada(s)."
        )
        if self.falhas:
            texto += f" {len(self.falhas)} item(ns) não gravado(s): veja o relatório."
        return texto


class _Codigos:
    """Próximo código numérico livre para cliente novo que veio sem código.

    Os códigos do banco são lidos uma vez só; os da planilha ficam reservados
    desde o começo, para o número gerado não roubar o de uma linha mais abaixo.
    """

    def __init__(self, sessao: Session, da_planilha: set[str]) -> None:
        self.usados = set(sessao.scalars(select(Cliente.codigo))) | da_planilha
        self.proximo = max((int(c) for c in self.usados if c.isdigit()), default=0) + 1

    def gerar(self) -> str:
        while str(self.proximo) in self.usados:
            self.proximo += 1
        codigo = str(self.proximo)
        self.usados.add(codigo)
        return codigo


def importar(sessao: Session, previa: Previa, atualizar_existentes: bool = True) -> Resultado:
    """Grava as linhas marcadas. Roda dentro da transação de quem chamou.

    A conferência é refeita aqui (o banco pode ter mudado desde a prévia) e cada
    item que o cadastro recusar vira linha no relatório, sem parar o resto.
    """
    resultado = Resultado()
    registro = Importacao(
        arquivo=str(previa.arquivo),
        aba=f"Personalizada: {previa.layout.nome or previa.aba}"[:60],
        observacoes=previa.resumo(),
    )
    sessao.add(registro)
    sessao.flush()
    resultado.importacao_id = registro.id

    def falhou(linha: int, campo: str, valor: str, motivo: str, gravidade: str = "ERRO"):
        problema = Problema(previa.aba, linha, campo, valor, motivo, gravidade)
        resultado.falhas.append(problema)
        sessao.add(ImportacaoErro(
            importacao_id=registro.id, linha=linha, aba=previa.aba[:60], campo=campo[:60],
            valor=valor[:500], gravidade=gravidade, motivo=motivo,
        ))

    conferir(sessao, previa, atualizar_existentes)
    for problema in previa.todos_problemas():
        sessao.add(ImportacaoErro(
            importacao_id=registro.id, linha=problema.linha, aba=problema.aba[:60],
            campo=problema.campo[:60], valor=problema.valor[:500],
            gravidade=problema.gravidade, motivo=problema.motivo,
        ))
    marcadas = previa.importaveis
    resultado.linhas_ignoradas = len(previa.linhas) - len(marcadas)
    # só as linhas marcadas contam: desmarcar uma linha tira os dados dela do cliente
    planos = _agrupar_clientes(previa.aba, marcadas, avisar=False)

    # --- clientes (o vínculo com empresa vem depois: a empresa pode ser nova)
    clientes: dict[str, Cliente] = {}
    criados: set[str] = set()
    codigos = _Codigos(sessao, {
        plano.valores["cliente_codigo"] for plano in planos.values()
        if plano.valores.get("cliente_codigo")
    })
    for chave, plano in planos.items():
        primeira = plano.linhas[0]
        existente, _ = _cliente_no_banco(sessao, plano)
        try:
            if existente is not None:
                clientes[chave] = existente
                if not atualizar_existentes:
                    resultado.clientes_sem_alteracao += 1
                    continue
                novo, mudou = _mesclar(_dados_de(existente), plano, None)
                if mudou:
                    juros = novo.juros_dia
                    servico_clientes.atualizar(sessao, existente, novo)
                    if "juros/dia" in mudou:
                        servico_clientes.registrar_juros(sessao, existente, juros, hoje(), MOTIVO)
                    resultado.clientes_atualizados += 1
                else:
                    resultado.clientes_sem_alteracao += 1
                continue
            valores = plano.valores
            codigo = valores.get("cliente_codigo") or codigos.gerar()
            dados, _ = _mesclar(
                servico_clientes.DadosCliente(
                    codigo=codigo, nome="",
                    tipo_pessoa=valores.get("cliente_tipo_pessoa") or TipoPessoa.PJ,
                    juros_dia=Decimal("0.00"), contatos=[],
                ),
                plano, None,
            )
            dados.observacoes = (dados.observacoes or MOTIVO).strip()
            clientes[chave] = servico_clientes.criar(sessao, dados)
            criados.add(chave)
            resultado.clientes_criados += 1
        except ErroDeNegocio as falha:
            falhou(primeira, "Cliente", plano.valores.get("cliente_nome")
                   or plano.valores.get("cliente_codigo", ""), str(falha))
    sessao.flush()

    # --- vínculo com a empresa e suspensão
    for chave, plano in planos.items():
        cliente = clientes.get(chave)
        if cliente is None:
            continue
        primeira = plano.linhas[0]
        criado_agora = chave in criados
        referencia = plano.valores.get("empresa_vinculo")
        if (referencia and cliente.tipo_pessoa is TipoPessoa.PF
                and (atualizar_existentes or criado_agora)):
            empresa, chave_empresa = _resolver_empresa(sessao, referencia, planos)
            empresa = empresa or clientes.get(chave_empresa or "")
            if empresa is None and chave_empresa:
                falhou(primeira, "Empresa que paga por ele", referencia,
                       "A empresa da planilha não foi gravada: vínculo não feito", "AVISO")
            if empresa is not None and empresa.id != cliente.empresa_id:
                dados = _dados_de(cliente)
                dados.empresa_id = empresa.id
                try:
                    servico_clientes.atualizar(sessao, cliente, dados)
                except ErroDeNegocio as falha:
                    falhou(primeira, "Empresa que paga por ele", referencia, str(falha), "AVISO")
        motivo = plano.valores.get("suspensao_motivo")
        if motivo and not cliente.cobranca_suspensa and (atualizar_existentes or criado_agora):
            try:
                servico_clientes.suspender_cobranca(
                    sessao, cliente, motivo, plano.valores.get("suspensao_ate")
                )
            except ErroDeNegocio as falha:
                falhou(primeira, "Suspender cobrança", motivo, str(falha), "AVISO")
    sessao.flush()

    # --- contratos e honorários, linha a linha
    contratos_feitos: set[tuple[int, int]] = set()
    for linha in marcadas:
        cliente = clientes.get(linha.chave_cliente)
        if cliente is None:
            if linha.tem("contrato_valor", "honorario_valor"):
                falhou(linha.linha, "Cliente", linha.rotulo_cliente,
                       "Cliente não foi gravado: contrato/honorário da linha não importado")
            continue
        if linha.tem("contrato_valor"):
            _gravar_contrato(sessao, linha, cliente, contratos_feitos, resultado, falhou,
                             atualizar_existentes)
        if linha.tem("honorario_valor"):
            _gravar_honorario(sessao, linha, cliente, registro.id, resultado, falhou)
    sessao.flush()

    registro.clientes_criados = resultado.clientes_criados
    registro.honorarios_criados = resultado.honorarios_criados
    registro.linhas_ignoradas = resultado.linhas_ignoradas
    sessao.flush()
    return resultado


def _gravar_contrato(sessao, linha, cliente, feitos, resultado, falhou, atualizar_existentes
                     ) -> None:
    valores = linha.valores
    tipo = valores["contrato_tipo"]
    if (cliente.id, tipo.id) in feitos:
        return
    feitos.add((cliente.id, tipo.id))
    existente = _contrato_existente(sessao, cliente.id, tipo.id)
    forma = valores.get("contrato_forma_envio")
    try:
        if existente is None:
            servico_contratos.criar(sessao, servico_contratos.DadosContrato(
                cliente_id=cliente.id,
                tipo_id=tipo.id,
                valor=valores["contrato_valor"],
                periodicidade=valores.get("contrato_periodicidade") or tipo.periodicidade,
                dia_vencimento=valores.get("contrato_dia_vencimento") or 10,
                meses_vencimento=valores.get("contrato_meses_vencimento", 1),
                meses_referencia=valores.get("contrato_meses_referencia") or "",
                forma_envio_id=forma.id if forma else None,
                meio=valores.get("contrato_meio") or MeioCobranca.BOLETO,
                vigencia_inicio=valores.get("contrato_vigencia_inicio"),
                vigencia_fim=valores.get("contrato_vigencia_fim"),
                suspenso=bool(valores.get("contrato_suspenso")),
                observacoes=valores.get("contrato_observacoes") or "",
            ))
            resultado.contratos_criados += 1
            return
        if not atualizar_existentes:
            return
        # primeiro o que valida (atualizar), depois o reajuste: se a validação
        # recusar, o valor não fica trocado pela metade
        mudou = False
        if _mudancas_contrato(existente, valores):
            servico_contratos.atualizar(sessao, existente, servico_contratos.DadosContrato(
                cliente_id=cliente.id,
                tipo_id=tipo.id,
                valor=existente.valor,
                periodicidade=valores.get("contrato_periodicidade") or existente.periodicidade,
                dia_vencimento=valores.get("contrato_dia_vencimento") or existente.dia_vencimento,
                meses_vencimento=valores.get("contrato_meses_vencimento",
                                             existente.meses_vencimento),
                meses_referencia=valores.get("contrato_meses_referencia",
                                             existente.meses_referencia),
                forma_envio_id=forma.id if forma else existente.forma_envio_id,
                meio=valores.get("contrato_meio") or existente.meio,
                vigencia_inicio=valores.get("contrato_vigencia_inicio")
                or existente.vigencia_inicio,
                vigencia_fim=valores.get("contrato_vigencia_fim", existente.vigencia_fim),
                suspenso=valores.get("contrato_suspenso", existente.suspenso),
                observacoes=valores.get("contrato_observacoes", existente.observacoes),
            ))
            mudou = True
        if valores["contrato_valor"] != existente.valor:
            servico_contratos.reajustar(sessao, existente, valores["contrato_valor"], hoje(),
                                        MOTIVO)
            mudou = True
        if mudou:
            resultado.contratos_atualizados += 1
    except ErroDeNegocio as falha:
        falhou(linha.linha, "Contrato", f"{linha.rotulo_cliente} – {tipo.nome}", str(falha))


def _gravar_honorario(sessao, linha, cliente, importacao_id, resultado, falhou) -> None:
    valores = linha.valores
    tipo = valores["honorario_tipo"]
    competencia = valores.get("honorario_competencia") or Competencia()
    status = valores.get("honorario_status") or StatusPagamento.EM_ABERTO
    if _honorario_existente(sessao, cliente.id, tipo.id, competencia,
                            valores["honorario_vencimento"], valores["honorario_valor"],
                            status):
        resultado.honorarios_existentes += 1
        return
    observacoes = [valores.get("honorario_observacoes") or ""]
    if status is StatusPagamento.PAGO:
        observacoes.append(
            "Importado como pago: o total foi congelado pelo valor do honorário."
        )
    forma = valores.get("honorario_forma_envio")
    envio = valores.get("honorario_situacao_envio") or SituacaoEnvio.NAO_ENVIADO
    try:
        honorario = servico_honorarios.criar(sessao, servico_honorarios.DadosHonorario(
            cliente_id=cliente.id,
            tipo_id=tipo.id,
            valor=valores["honorario_valor"],
            vencimento=valores["honorario_vencimento"],
            competencia=competencia,
            descricao=valores.get("honorario_descricao") or "",
            meio=valores.get("honorario_meio") or MeioCobranca.BOLETO,
            forma_envio_id=forma.id if forma else None,
            observacoes="\n".join(texto for texto in observacoes if texto),
            numero_boleto=valores.get("honorario_numero_boleto") or "",
            situacao_envio=envio,
            enviado_em=valores["honorario_vencimento"] if envio is SituacaoEnvio.ENVIADO
            else None,
        ))
    except ErroDeNegocio as falha:
        falhou(linha.linha, "Honorário", f"{linha.rotulo_cliente} {competencia.rotulo()}",
               str(falha))
        return
    honorario.origem_importacao_id = importacao_id
    honorario.status = status
    if status is StatusPagamento.PAGO:
        honorario.saldo_aberto = ZERO
        honorario.total_congelado = honorario.valor_original
        honorario.pago_em = valores.get("honorario_pago_em")
    elif status in (StatusPagamento.CANCELADO, StatusPagamento.RENEGOCIADO):
        honorario.saldo_aberto = ZERO
    resultado.honorarios_criados += 1
