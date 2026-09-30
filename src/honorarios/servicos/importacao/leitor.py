"""Leitura da planilha antiga, com normalização e relatório de problemas.

O que a planilha tem de irregular está mapeado aqui: competência em formatos
misturados (texto, data, serial do Excel, nome de serviço), linhas sem
vencimento ou sem valor, o mesmo cliente com códigos diferentes e o juros ao dia
divergindo entre as linhas de um mesmo cliente. Nada é adivinhado em silêncio:
cada ajuste entra no relatório como AVISO e cada linha impossível como ERRO.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from ...dominio.enums import (
    FormaPagamento,
    MeioCobranca,
    SituacaoEnvio,
    StatusPagamento,
    SubtipoCliente,
    TipoPessoa,
)
from ...dominio.valores import (
    ZERO,
    Competencia,
    dinheiro,
    normalizar_busca,
    so_digitos,
    validar_documento,
)

ABA_BASE = "Controle Financeiro"
ABA_IRPF = "Controle de IRPF"
ABA_RECALL = "Recall Inadimplentes"

# Status da planilha -> (status, forma de pagamento, situação de envio)
MAPA_STATUS = {
    "em aberto": (StatusPagamento.EM_ABERTO, None, None),
    "pago": (StatusPagamento.PAGO, FormaPagamento.PIX, None),
    "pago direto para eloisa": (StatusPagamento.PAGO, FormaPagamento.DIRETO_ELOISA, None),
    "cancelado": (StatusPagamento.CANCELADO, None, None),
    "renegociado": (StatusPagamento.RENEGOCIADO, None, None),
    "enviado para cliente": (StatusPagamento.EM_ABERTO, None, SituacaoEnvio.ENVIADO),
}

# Tipo de faturamento da planilha -> nome do tipo de honorário no sistema
MAPA_TIPO = {
    "mensal": "Honorário mensal",
    "trimestral": "Honorário trimestral",
    "anual": "Honorário anual",
    "imposto de renda": "Imposto de Renda",
    "ata de socios": "Ata de Sócios",
    "alteracao contratual": "Alteração contratual",
    "ganho de capital": "Ganho de Capital",
    "suspenso temporariamente": "Honorário mensal",
}

MAPA_PESSOA = {
    "pf": (TipoPessoa.PF, SubtipoCliente.NENHUM),
    "pj": (TipoPessoa.PJ, SubtipoCliente.NENHUM),
    "domestica": (TipoPessoa.PF, SubtipoCliente.DOMESTICA),
    "imposto de renda": (TipoPessoa.PF, SubtipoCliente.NENHUM),
    "irpf": (TipoPessoa.PF, SubtipoCliente.NENHUM),
}


@dataclass
class Problema:
    aba: str
    linha: int
    campo: str
    valor: str
    motivo: str
    gravidade: str = "AVISO"

    @property
    def bloqueia(self) -> bool:
        return self.gravidade == "ERRO"


@dataclass
class CobrancaLida:
    quando: date
    texto: str


@dataclass
class LinhaLida:
    aba: str
    linha: int
    codigo: str
    nome: str
    documento: str
    tipo_pessoa: TipoPessoa
    subtipo: SubtipoCliente
    tipo_honorario: str
    competencia: Competencia
    descricao: str
    vencimento: date | None
    valor: Decimal
    juros_dia: Decimal | None
    status: StatusPagamento
    forma_pagamento: FormaPagamento | None
    situacao_envio: SituacaoEnvio
    forma_envio: str
    meio: MeioCobranca
    observacoes: str = ""
    cobrancas: list[CobrancaLida] = field(default_factory=list)
    problemas: list[Problema] = field(default_factory=list)
    importar: bool = True
    # chave do cliente já resolvida pela deduplicação (documento, código ou nome)
    chave_cliente: str = ""

    @property
    def bloqueada(self) -> bool:
        return any(problema.bloqueia for problema in self.problemas)


@dataclass
class ClienteLido:
    codigo: str
    nome: str
    documento: str
    tipo_pessoa: TipoPessoa
    subtipo: SubtipoCliente
    juros_dia: Decimal | None = None
    jurosdivergentes: list[Decimal] = field(default_factory=list)
    origem: str = ABA_BASE

    @property
    def chave(self) -> str:
        return self.documento or f"codigo:{self.codigo}"


@dataclass
class Leitura:
    arquivo: Path
    clientes: dict[str, ClienteLido] = field(default_factory=dict)
    linhas: list[LinhaLida] = field(default_factory=list)
    problemas: list[Problema] = field(default_factory=list)
    # índices de deduplicação: o mesmo cliente aparece nas três abas com
    # identificação diferente (a de recall não tem CPF/CNPJ, a de IRPF usa outro código)
    por_codigo: dict[str, str] = field(default_factory=dict)
    por_nome: dict[str, str] = field(default_factory=dict)

    @property
    def importaveis(self) -> list[LinhaLida]:
        return [linha for linha in self.linhas if linha.importar and not linha.bloqueada]

    @property
    def total_cobrancas(self) -> int:
        return sum(len(linha.cobrancas) for linha in self.linhas)

    def resumo(self) -> str:
        from ...dominio.valores import formatar_moeda

        total = sum((linha.valor for linha in self.importaveis), start=ZERO)
        erros = sum(1 for p in self.todos_problemas() if p.bloqueia)
        return (
            f"{len(self.clientes)} cliente(s), {len(self.importaveis)} honorário(s) "
            f"({formatar_moeda(total)}) e {self.total_cobrancas} cobrança(s) a importar. "
            f"{erros} linha(s) com erro."
        )

    def todos_problemas(self) -> list[Problema]:
        return [*self.problemas, *(p for linha in self.linhas for p in linha.problemas)]


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _data(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        return date.fromordinal(date(1899, 12, 30).toordinal() + int(valor))
    from ...dominio.valores import ler_data

    try:
        return ler_data(str(valor))
    except ValueError:
        return None


def _decimal(valor) -> Decimal | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, str):
        from ...dominio.valores import ler_moeda

        try:
            return ler_moeda(valor)
        except ValueError:
            return None
    try:
        return dinheiro(valor)
    except ValueError:
        return None


def _juros(valor) -> Decimal | None:
    if valor is None or valor == "":
        return None
    try:
        return Decimal(str(valor).replace(",", ".")).quantize(Decimal("0.0001"))
    except Exception:  # noqa: BLE001
        return None


def separar_cobrancas(
    texto: str, aba: str, linha: int
) -> tuple[list[CobrancaLida], list[Problema], str]:
    """Quebra a coluna Observações em cobranças datadas.

    Linhas no formato ``dd/mm/aaaa - texto`` viram cobrança; o que não casa fica
    como observação do honorário e entra no relatório.
    """
    import re

    padrao = re.compile(r"^\s*(\d{2}/\d{2}/\d{4})\s*[-–:]\s*(.+)$")
    cobrancas: list[CobrancaLida] = []
    problemas: list[Problema] = []
    sobras: list[str] = []
    for bruto in (texto or "").replace("\r\n", "\n").split("\n"):
        conteudo = bruto.strip()
        if not conteudo:
            continue
        casamento = padrao.match(conteudo)
        if not casamento:
            sobras.append(conteudo)
            problemas.append(
                Problema(
                    aba,
                    linha,
                    "Observações",
                    conteudo[:120],
                    "Sem data no início: mantido como observação, não virou cobrança",
                )
            )
            continue
        from ...dominio.valores import ler_data

        try:
            quando = ler_data(casamento.group(1))
        except ValueError:
            sobras.append(conteudo)
            problemas.append(
                Problema(aba, linha, "Observações", conteudo[:120], "Data inválida na cobrança")
            )
            continue
        cobrancas.append(CobrancaLida(quando=quando, texto=casamento.group(2).strip()))
    return cobrancas, problemas, "\n".join(sobras)


def _pessoa(tipo_texto: str) -> tuple[TipoPessoa, SubtipoCliente]:
    return MAPA_PESSOA.get(normalizar_busca(tipo_texto), (TipoPessoa.PJ, SubtipoCliente.NENHUM))


def _tipo_honorario(tipo_faturamento: str, competencia_bruta: str, aba: str) -> tuple[str, bool]:
    """Devolve o nome do tipo e se a competência deve ser descartada."""
    chave = normalizar_busca(tipo_faturamento)
    if aba == ABA_IRPF:
        return "Imposto de Renda", False
    alvo = normalizar_busca(competencia_bruta)
    for nome_chave, nome in MAPA_TIPO.items():
        if alvo and alvo == nome_chave:
            # a competência guardava o nome do serviço (ex.: "ATA DE SOCIOS")
            return nome, True
    return MAPA_TIPO.get(chave, "Honorário mensal"), False


def ler(arquivo: Path | str) -> Leitura:
    """Lê as três abas úteis da planilha e devolve tudo normalizado."""
    from openpyxl import load_workbook

    caminho = Path(arquivo)
    leitura = Leitura(arquivo=caminho)
    livro = load_workbook(caminho, data_only=True, read_only=True)

    if ABA_RECALL in livro.sheetnames:
        _ler_recall(livro[ABA_RECALL], leitura)
    if ABA_BASE in livro.sheetnames:
        _ler_base(livro[ABA_BASE], leitura)
    else:
        leitura.problemas.append(
            Problema("", 0, "arquivo", caminho.name, f"Aba {ABA_BASE} não encontrada", "ERRO")
        )
    if ABA_IRPF in livro.sheetnames:
        _ler_irpf(livro[ABA_IRPF], leitura)
    livro.close()
    _consolidar_juros(leitura)
    return leitura


def _registrar_cliente(
    leitura: Leitura,
    codigo: str,
    nome: str,
    documento: str,
    tipo_pessoa: TipoPessoa,
    subtipo: SubtipoCliente,
    juros: Decimal | None,
    aba: str,
    linha: int,
) -> tuple[str, ClienteLido]:
    documento = so_digitos(documento)
    if documento and not validar_documento(documento):
        leitura.problemas.append(
            Problema(
                aba, linha, "CPF/CNPJ", documento, "Dígito verificador inválido: gravado vazio"
            )
        )
        documento = ""

    if documento:
        # o documento decide PF/PJ: a coluna Tipo da planilha traz "Imposto de
        # Renda" em linhas de empresa, o que viraria PF sem CPF e quebraria o cadastro
        pelo_documento = TipoPessoa.PF if len(documento) == 11 else TipoPessoa.PJ
        if pelo_documento is not tipo_pessoa:
            leitura.problemas.append(
                Problema(
                    aba,
                    linha,
                    "Tipo",
                    f"{tipo_pessoa.value} com documento de {pelo_documento.value}",
                    f"{nome}: tipo corrigido para {pelo_documento.rotulo} pelo CPF/CNPJ",
                )
            )
        tipo_pessoa = pelo_documento
        if tipo_pessoa is TipoPessoa.PJ:
            subtipo = SubtipoCliente.NENHUM

    alvo_nome = normalizar_busca(nome)
    chave = (
        (documento and leitura.clientes.get(documento) and documento)
        or leitura.por_codigo.get(codigo)
        or leitura.por_nome.get(alvo_nome)
        or documento
        or f"codigo:{codigo}"
        or f"nome:{alvo_nome}"
    )

    existente = leitura.clientes.get(chave)
    if existente is None:
        existente = ClienteLido(
            codigo=codigo,
            nome=nome,
            documento=documento,
            tipo_pessoa=tipo_pessoa,
            subtipo=subtipo,
            juros_dia=juros,
            origem=aba,
        )
        leitura.clientes[chave] = existente
    else:
        if not existente.documento and documento:
            existente.documento = documento
        if not existente.codigo and codigo:
            existente.codigo = codigo
        if documento:
            existente.tipo_pessoa = tipo_pessoa
            if subtipo is not SubtipoCliente.NENHUM:
                existente.subtipo = subtipo
        elif existente.origem == ABA_RECALL and aba != ABA_RECALL:
            existente.tipo_pessoa = tipo_pessoa
            existente.subtipo = subtipo
            existente.origem = aba
        if existente.codigo and codigo and existente.codigo != codigo:
            leitura.problemas.append(
                Problema(
                    aba,
                    linha,
                    "Código Cliente",
                    f"{existente.codigo} / {codigo}",
                    f"{nome} aparece com dois códigos: mantido {existente.codigo}",
                )
            )
    if codigo:
        leitura.por_codigo.setdefault(codigo, chave)
    if alvo_nome:
        leitura.por_nome.setdefault(alvo_nome, chave)
    if juros is not None and juros not in existente.jurosdivergentes:
        existente.jurosdivergentes.append(juros)
    return chave, existente


def _consolidar_juros(leitura: Leitura) -> None:
    """Quando o cliente tem juros/dia diferente entre linhas, fica o maior e avisa."""
    for cliente in leitura.clientes.values():
        valores = [v for v in cliente.jurosdivergentes if v is not None]
        if not valores:
            continue
        cliente.juros_dia = max(valores)
        if len(set(valores)) > 1:
            leitura.problemas.append(
                Problema(
                    cliente.origem,
                    0,
                    "Juros ao Dia",
                    " / ".join(str(v.normalize()) for v in sorted(set(valores))),
                    f"{cliente.nome}: valores divergentes na planilha. Sugerido "
                    f"{cliente.juros_dia.normalize()} — confira antes de confirmar",
                )
            )


def _ler_base(aba, leitura: Leitura) -> None:
    """Aba principal: cabeçalho na linha 2, dados a partir da 3."""
    nome_aba = ABA_BASE
    for numero, valores in enumerate(aba.iter_rows(min_row=3, values_only=True), start=3):
        if valores is None or all(valor in (None, "") for valor in valores):
            continue
        (
            competencia_bruta,
            codigo,
            nome,
            documento,
            status_bruto,
            vencimento,
            valor,
            juros,
            *_resto,
        ) = list(valores) + [None] * (16 - len(valores))
        observacoes = valores[12] if len(valores) > 12 else None
        tipo_bruto = valores[13] if len(valores) > 13 else None
        faturamento = valores[14] if len(valores) > 14 else None
        forma_envio = valores[15] if len(valores) > 15 else None

        linha = _montar_linha(
            leitura=leitura,
            aba=nome_aba,
            numero=numero,
            competencia_valor=competencia_bruta,
            competencia_bruta=_texto(competencia_bruta),
            codigo=_texto(codigo),
            nome=_texto(nome),
            documento=_texto(documento),
            status_bruto=_texto(status_bruto),
            vencimento=vencimento,
            valor=valor,
            juros=juros,
            tipo_bruto=_texto(tipo_bruto),
            faturamento=_texto(faturamento),
            forma_envio=_texto(forma_envio),
            observacoes=_texto(observacoes),
            meio=MeioCobranca.BOLETO,
        )
        if linha is not None:
            leitura.linhas.append(linha)


def _ler_irpf(aba, leitura: Leitura) -> None:
    """Aba de IRPF: mesma estrutura, colunas em outra ordem."""
    nome_aba = ABA_IRPF
    for numero, valores in enumerate(aba.iter_rows(min_row=3, values_only=True), start=3):
        if valores is None or all(valor in (None, "") for valor in valores):
            continue
        dados = list(valores) + [None] * (19 - len(valores))
        (
            tipo_bruto,
            faturamento,
            competencia_bruta,
            codigo,
            nome,
            documento,
            vencimento,
            valor,
            juros,
            forma_envio,
            boleto_ou_pix,
            status_bruto,
            *_resto,
        ) = dados
        observacoes = dados[17]
        meio = (
            MeioCobranca.PIX
            if normalizar_busca(_texto(boleto_ou_pix)) == "pix"
            else MeioCobranca.BOLETO
        )
        linha = _montar_linha(
            leitura=leitura,
            aba=nome_aba,
            numero=numero,
            competencia_valor=competencia_bruta,
            competencia_bruta=_texto(competencia_bruta),
            codigo=_texto(codigo),
            nome=_texto(nome),
            documento=_texto(documento),
            status_bruto=_texto(status_bruto),
            vencimento=vencimento,
            valor=valor,
            juros=juros,
            tipo_bruto=_texto(tipo_bruto),
            faturamento=_texto(faturamento) or "Imposto de Renda",
            forma_envio=_texto(forma_envio),
            observacoes=_texto(observacoes),
            meio=meio,
        )
        if linha is not None:
            leitura.linhas.append(linha)


def _ler_recall(aba, leitura: Leitura) -> None:
    """A aba de recall é a lista mestra de clientes: só semeia o cadastro."""
    for numero, valores in enumerate(aba.iter_rows(min_row=5, values_only=True), start=5):
        if valores is None or all(valor in (None, "") for valor in valores):
            continue
        codigo = _texto(valores[0])
        nome = _texto(valores[1])
        tipo_bruto = _texto(valores[2]) if len(valores) > 2 else ""
        if not nome:
            continue
        tipo_pessoa, subtipo = _pessoa(tipo_bruto)
        _registrar_cliente(
            leitura, codigo, nome, "", tipo_pessoa, subtipo, None, ABA_RECALL, numero
        )


def _montar_linha(
    leitura: Leitura,
    aba: str,
    numero: int,
    competencia_valor,
    competencia_bruta: str,
    codigo: str,
    nome: str,
    documento: str,
    status_bruto: str,
    vencimento,
    valor,
    juros,
    tipo_bruto: str,
    faturamento: str,
    forma_envio: str,
    observacoes: str,
    meio: MeioCobranca,
) -> LinhaLida | None:
    problemas: list[Problema] = []
    if not nome:
        leitura.problemas.append(
            Problema(aba, numero, "Cliente", "", "Linha sem nome de cliente: ignorada", "ERRO")
        )
        return None

    tipo_pessoa, subtipo = _pessoa(tipo_bruto)
    juros_lido = _juros(juros)
    chave_cliente, cliente = _registrar_cliente(
        leitura, codigo, nome, documento, tipo_pessoa, subtipo, juros_lido, aba, numero
    )

    nome_tipo, descartar_competencia = _tipo_honorario(faturamento, competencia_bruta, aba)
    competencia = Competencia()
    descricao = ""
    if descartar_competencia:
        descricao = nome_tipo
        problemas.append(
            Problema(
                aba,
                numero,
                "Competência",
                competencia_bruta,
                f"Continha o nome do serviço: virou o tipo {nome_tipo}, sem competência",
            )
        )
    else:
        try:
            competencia = Competencia.de_texto(
                competencia_valor if competencia_valor not in (None, "") else None
            )
        except ValueError:
            problemas.append(
                Problema(aba, numero, "Competência", competencia_bruta, "Formato não reconhecido")
            )
        if competencia.vazia and competencia_bruta:
            problemas.append(
                Problema(
                    aba, numero, "Competência", competencia_bruta,
                    "Não é competência: gravado como descrição",
                )
            )
            descricao = competencia_bruta

    data_vencimento = _data(vencimento)
    if data_vencimento is None:
        problemas.append(
            Problema(
                aba, numero, "Vencimento", _texto(vencimento),
                "Sem vencimento: não é possível calcular atraso", "ERRO",
            )
        )
    valor_lido = _decimal(valor)
    if valor_lido is None or valor_lido <= 0:
        problemas.append(
            Problema(
                aba, numero, "Valor do Boleto", _texto(valor),
                "Valor ausente ou zerado", "ERRO",
            )
        )
    if juros_lido is None:
        problemas.append(
            Problema(
                aba, numero, "Juros ao Dia", _texto(juros),
                f"Sem juros ao dia na planilha: usará o do cadastro de {nome}",
            )
        )

    status, forma_pagamento, envio = MAPA_STATUS.get(
        normalizar_busca(status_bruto), (StatusPagamento.EM_ABERTO, None, None)
    )
    if status_bruto and normalizar_busca(status_bruto) not in MAPA_STATUS:
        problemas.append(
            Problema(
                aba, numero, "Status Pagamento", status_bruto,
                "Status desconhecido: importado como Em aberto",
            )
        )

    cobrancas, problemas_obs, sobra = separar_cobrancas(observacoes, aba, numero)
    problemas.extend(problemas_obs)

    if competencia.ano is None and competencia.mes is None and not descricao and data_vencimento:
        descricao = descricao or nome_tipo

    return LinhaLida(
        aba=aba,
        linha=numero,
        codigo=cliente.codigo,
        nome=nome,
        documento=cliente.documento,
        tipo_pessoa=tipo_pessoa,
        subtipo=subtipo,
        tipo_honorario=nome_tipo,
        competencia=competencia,
        descricao=descricao,
        vencimento=data_vencimento,
        valor=valor_lido or ZERO,
        juros_dia=juros_lido,
        status=status,
        forma_pagamento=forma_pagamento,
        situacao_envio=envio or SituacaoEnvio.NAO_ENVIADO,
        forma_envio=forma_envio,
        meio=meio,
        observacoes=sobra,
        cobrancas=cobrancas,
        problemas=problemas,
        chave_cliente=chave_cliente,
    )
