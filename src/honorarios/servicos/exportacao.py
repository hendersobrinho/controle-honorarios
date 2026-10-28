"""Exportações: planilha Excel da lista filtrada e PDF do extrato de débitos."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config.caminhos import pasta_exportacoes
from ..dominio.valores import ZERO, formatar_data, formatar_moeda, hoje
from .honorarios import LinhaHonorario

FORMATO_MOEDA = 'R$ #,##0.00'
CABECALHO = (
    "Cliente",
    "Código",
    "CPF/CNPJ",
    "Tipo",
    "Grupo",
    "Competência",
    "Vencimento",
    "Dias em atraso",
    "Valor original",
    "Juros",
    "Total atualizado",
    "Status",
    "Cobranças",
    "Última cobrança",
    "Promessa",
)


def _nome_arquivo(prefixo: str, extensao: str) -> Path:
    marca = hoje().strftime("%Y%m%d")
    destino = pasta_exportacoes() / f"{prefixo}-{marca}.{extensao}"
    contador = 2
    while destino.exists():
        destino = pasta_exportacoes() / f"{prefixo}-{marca}-{contador}.{extensao}"
        contador += 1
    return destino


def para_excel(
    linhas: Sequence[LinhaHonorario],
    destino: Path | None = None,
    titulo: str = "Honorários",
) -> Path:
    """Planilha da lista já filtrada, com os valores atualizados e um total."""
    arquivo = Path(destino) if destino else _nome_arquivo("honorarios", "xlsx")
    livro = Workbook()
    aba = livro.active
    aba.title = titulo[:31]

    aba["A1"] = titulo
    aba["A1"].font = Font(bold=True, size=14)
    aba["A2"] = f"Gerado em {formatar_data(hoje())} · {len(linhas)} lançamento(s)"
    aba["A2"].font = Font(italic=True, color="666666")

    for coluna, nome in enumerate(CABECALHO, start=1):
        celula = aba.cell(row=4, column=coluna, value=nome)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="EEF1F5")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for indice, linha in enumerate(linhas, start=5):
        honorario = linha.honorario
        valores = (
            linha.cliente.nome,
            linha.cliente.codigo,
            linha.cliente.documento,
            honorario.tipo.nome,
            linha.grupo.nome,
            honorario.competencia_rotulo,
            honorario.vencimento,
            linha.atualizacao.dias_atraso,
            float(linha.atualizacao.valor_original),
            float(linha.atualizacao.juros),
            float(linha.atualizacao.total),
            honorario.status.rotulo,
            linha.cobrancas,
            linha.ultima_cobranca,
            linha.promessa,
        )
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=indice, column=coluna, value=valor)
            if coluna in (9, 10, 11):
                celula.number_format = FORMATO_MOEDA
            if coluna in (7, 14, 15):
                celula.number_format = "DD/MM/YYYY"

    ultima = 4 + len(linhas) + 1
    aba.cell(row=ultima, column=8, value="Total").font = Font(bold=True)
    for coluna in (9, 10, 11):
        letra = get_column_letter(coluna)
        celula = aba.cell(row=ultima, column=coluna)
        celula.value = f"=SUM({letra}5:{letra}{4 + len(linhas)})" if linhas else 0
        celula.number_format = FORMATO_MOEDA
        celula.font = Font(bold=True)

    larguras = (34, 8, 20, 22, 20, 13, 12, 8, 14, 12, 16, 14, 10, 14, 12)
    for coluna, largura in enumerate(larguras, start=1):
        aba.column_dimensions[get_column_letter(coluna)].width = largura
    aba.freeze_panes = "A5"

    livro.save(arquivo)
    return arquivo


def extrato_pdf(
    cliente,
    linhas: Sequence[LinhaHonorario],
    recebedor,
    destino: Path | None = None,
    data_referencia: date | None = None,
) -> Path:
    """Extrato de débitos da empresa, para enviar ao cliente (no molde dos PDFs)."""
    from reportlab.platypus import Spacer

    from .pdf_relatorio import RelatorioPdf, texto

    referencia = data_referencia or hoje()
    arquivo = Path(destino) if destino else _nome_arquivo(
        f"extrato-{cliente.codigo}", "pdf"
    )
    r = RelatorioPdf(
        "Extrato de débitos em aberto",
        f"{cliente.codigo} – {cliente.nome}",
        recebedor,
        referencia,
        rodape="Extrato de débitos",
    )
    total = sum((linha.atualizacao.total for linha in linhas), start=ZERO)
    dados = [["Honorário", "Competência", "Vencimento", "Atraso", "Valor", "Juros", "Total"]]
    for linha in linhas:
        dados.append([
            r.p(texto(linha.honorario.tipo.nome), "celula"),
            linha.honorario.competencia_rotulo,
            formatar_data(linha.honorario.vencimento),
            f"{linha.atualizacao.dias_atraso} d",
            formatar_moeda(linha.atualizacao.valor_original),
            formatar_moeda(linha.atualizacao.juros),
            formatar_moeda(linha.atualizacao.total),
        ])
    dados.append(["", "", "", "", "", "Total", formatar_moeda(total)])
    corpo = [
        *r.inicio(),
        r.p(f"<b>Cliente:</b> {texto(cliente.nome)} (código {cliente.codigo})"),
        r.p(f"<b>Valores atualizados para:</b> {formatar_data(referencia)}"),
        Spacer(1, 10),
        r.tabela(dados, r.larguras((0.28, 0.12, 0.12, 0.08, 0.14, 0.12, 0.14)), total=True),
        Spacer(1, 12),
        r.p(f"Pagamento por PIX – {texto(recebedor.razao_social)}<br/>"
            f"{texto(recebedor.rotulo_chave)}: {texto(recebedor.chave)}"),
        Spacer(1, 6),
        r.p("Juros de mora calculados por dia de atraso, conforme contrato. "
            "Após o pagamento, por gentileza, encaminhe o comprovante.", "nota"),
    ]
    return r.salvar(corpo, arquivo, f"Extrato de débitos – {cliente.nome}")
