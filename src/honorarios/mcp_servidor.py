"""Servidor MCP do Financeiro: o Claude trabalha nos honorários como o escritório.

Roda por stdio (``honorarios-mcp``) e usa a mesma camada de serviços das telas,
então toda regra de negócio vale igual: juros, validações, log de alterações. Não
há modo somente leitura — as ferramentas leem e gravam tudo (clientes, honorários,
pagamentos, cobranças, contratos, configurações). Cada ferramenta grava numa
transação própria: se uma regra recusa, nada fica pela metade.

Ferramentas que apagam ou desfazem têm ``destructiveHint`` para o cliente MCP
poder distingui-las; isso é só dica, não bloqueio.
"""

from __future__ import annotations

import functools
import json
import sys
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import fields, is_dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from enum import Enum

from mcp.server.fastmcp import FastMCP
from mcp.types import ToolAnnotations
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from .dominio.enums import (
    Canal,
    FormaPagamento,
    MeioCobranca,
    Periodicidade,
    ResultadoCobranca,
    SituacaoCliente,
    StatusPagamento,
    SubtipoCliente,
    TipoPessoa,
)
from .dominio.valores import Competencia, agora_utc, formatar_moeda, hoje, ler_data
from .infra import db, esquema
from .infra.modelos import Cobranca, Honorario, Pagamento
from .servicos import (
    calculo,
    clientes,
    cobrancas,
    conexao_claude,
    configuracoes,
    contratos,
    pagamentos,
    rankings,
    tipos,
)
from .servicos import honorarios as servico
from .servicos.erros import ErroDeNegocio

mcp = FastMCP(
    "financeiro-honorarios",
    instructions=(
        "Controle de honorários do escritório (clientes, honorários, juros de atraso, "
        "pagamentos, cobranças e contratos). Datas em AAAA-MM-DD ou DD/MM/AAAA; valores "
        "em reais como texto ('1234.56' ou '1.234,56'). Consulte antes de gravar e "
        "confirme com o usuário antes de excluir, cancelar ou estornar.\n"
        "E-mail do financeiro: verificar_emails → ler_email (já confere com o sistema) → "
        "baixar_quem_saiu_da_lista (confirme antes) → responder_financeiro (prefira "
        "modo='rascunho'; 'enviar' manda pela conta do sistema e não volta atrás).\n"
        "WhatsApp: cobrar_por_whatsapp abre a conversa com a mensagem e a chave Pix e "
        "copia a imagem; quem envia é o usuário. Registre a cobrança (registrar=True ou "
        "registrar_cobranca) só depois que ele confirmar o envio.\n"
        "Lista de contatos das empresas: ler_planilha para ver a lista → "
        "cadastrar_contatos_em_lote (gravar=false mostra o cruzamento; confira e repita "
        "com gravar=true).\n"
        "Planilhas: campos_de_importacao → salvar_layout_de_importacao → "
        "previa_de_importacao → importar_planilha.\n"
        "Banco: esquema_do_banco e consultar_banco (só leitura) para o que as ferramentas "
        "não cobrem; executar_no_banco grava por fora das regras do sistema — último "
        "recurso, sempre com confirmação."
    ),
)

LEITURA = ToolAnnotations(readOnlyHint=True)
GRAVA = ToolAnnotations(readOnlyHint=False, destructiveHint=False)
APAGA = ToolAnnotations(readOnlyHint=False, destructiveHint=True)

_preparado = False

# o que a tela de Configurações → Claude lê para dizer se o Claude está usando
CHAVE_ULTIMO_USO = conexao_claude.CHAVE_ULTIMO_USO
CHAVE_ULTIMA_FERRAMENTA = conexao_claude.CHAVE_ULTIMA_FERRAMENTA


def _anotar_uso(nome: str) -> None:
    """Grava quando e qual ferramenta o Claude chamou. Falhar aqui não derruba a chamada."""
    try:
        with db.transacao() as sessao:
            configuracoes.gravar(sessao, CHAVE_ULTIMO_USO, agora_utc().isoformat())
            configuracoes.gravar(sessao, CHAVE_ULTIMA_FERRAMENTA, nome)
    except Exception as falha:  # noqa: BLE001
        print(f"Não consegui anotar o uso de {nome}: {falha}", file=sys.stderr)


def _ferramenta(anotacoes: ToolAnnotations) -> Callable:
    """``mcp.tool`` que também anota o uso (o FastMCP lê a assinatura pelo ``wraps``)."""
    def decorar(funcao: Callable) -> Callable:
        @functools.wraps(funcao)
        def chamada(*args, **kwargs):
            try:
                return funcao(*args, **kwargs)
            finally:
                _anotar_uso(funcao.__name__)

        mcp.tool(annotations=anotacoes)(chamada)
        return chamada
    return decorar


# ----------------------------------------------------------------- infraestrutura
@contextmanager
def _sessao(gravar: bool = False) -> Iterator[Session]:
    """Sessão por chamada; ``gravar`` faz commit no fim e rollback no erro."""
    global _preparado
    if not _preparado:
        esquema.preparar()
        _preparado = True
    aberta = db.sessao()
    try:
        yield aberta
        if gravar:
            aberta.commit()
    except ErroDeNegocio as falha:
        aberta.rollback()
        raise ValueError(str(falha)) from falha
    except Exception:
        aberta.rollback()
        raise
    finally:
        aberta.close()


def _data(texto: str | None) -> date | None:
    return ler_data(texto) if texto else None


def _valor(texto: str | float | int | Decimal) -> Decimal:
    if isinstance(texto, str) and "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(str(texto).strip())
    except InvalidOperation as falha:
        raise ValueError(f"Valor inválido: {texto!r}.") from falha


def _enum(enumeracao, texto: str | None, padrao=None):
    if not texto:
        return padrao
    chave = texto.strip().upper().replace(" ", "_")
    try:
        return enumeracao(chave)
    except ValueError as falha:
        validos = ", ".join(item.value for item in enumeracao)
        raise ValueError(
            f"{texto!r} não vale para {enumeracao.__name__}: use {validos}."
        ) from falha


def _plano(valor):
    """Converte para algo que vira JSON: Decimal em texto, data em ISO, enum em valor."""
    if isinstance(valor, Decimal):
        return str(valor)
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    if isinstance(valor, Enum):
        return valor.value
    if is_dataclass(valor) and not isinstance(valor, type):
        return {campo.name: _plano(getattr(valor, campo.name)) for campo in fields(valor)}
    if isinstance(valor, dict):
        return {chave: _plano(item) for chave, item in valor.items()}
    if isinstance(valor, (list, tuple, set)):
        return [_plano(item) for item in valor]
    return valor


def _cliente(sessao: Session, cliente_id: int):
    encontrado = clientes.obter(sessao, cliente_id)
    if encontrado is None or encontrado.excluido_em is not None:
        raise ValueError(f"Cliente {cliente_id} não encontrado.")
    return encontrado


def _honorario(sessao: Session, honorario_id: int) -> Honorario:
    encontrado = sessao.get(Honorario, honorario_id)
    if encontrado is None:
        raise ValueError(f"Honorário {honorario_id} não encontrado.")
    return encontrado


def _linha(sessao: Session, honorario_id: int, referencia: date):
    h = _honorario(sessao, honorario_id)
    return servico.montar_linha(h, calculo.atualizar(h, referencia), referencia)


def _dados_cliente(c) -> dict:
    return {
        "id": c.id,
        "codigo": c.codigo,
        "nome": c.nome,
        "documento": c.documento,
        "tipo_pessoa": c.tipo_pessoa.value,
        "situacao": c.situacao.value,
        "empresa_id": c.empresa_id,
        "juros_dia": _plano(clientes.juros_atual(c)),
        "cobranca_suspensa": c.suspensao_ativa(),
        "motivo_suspensao": c.motivo_suspensao or None,
        "suspensa_ate": _plano(c.suspensa_ate),
        "observacoes": c.observacoes,
    }


def _dados_linha(linha) -> dict:
    h, a = linha.honorario, linha.atualizacao
    return {
        "id": h.id,
        "cliente_id": linha.cliente.id,
        "cliente": linha.cliente.nome,
        "tipo": h.tipo.nome,
        "competencia": h.competencia_rotulo,
        "descricao": h.descricao,
        "vencimento": _plano(h.vencimento),
        "status": h.status.value,
        "valor_original": _plano(a.valor_original),
        "saldo": _plano(a.saldo),
        "juros": _plano(a.juros),
        "total_atualizado": _plano(a.total),
        "dias_atraso": a.dias_atraso,
        "cobrancas": linha.cobrancas,
        "ultima_cobranca": _plano(linha.ultima_cobranca),
        "promessa": _plano(linha.promessa),
        "pago_em": _plano(h.pago_em),
        "numero_boleto": h.numero_boleto,
        "observacoes": h.observacoes,
    }


def _dados_pagamento(p: Pagamento) -> dict:
    return {
        "id": p.id,
        "honorario_id": p.honorario_id,
        "data": _plano(p.data),
        "valor_pago": _plano(p.valor_pago),
        "juros_cobrados": _plano(p.juros_cobrados),
        "desconto": _plano(p.desconto),
        "dias_atraso": p.dias_atraso,
        "forma_pagamento": p.forma_pagamento.value,
        "ativo": p.ativo,
    }


# --------------------------------------------------------------------- consultas
@_ferramenta(LEITURA)
def resumo_geral(data_referencia: str | None = None) -> dict:
    """Visão geral: quantos honorários atrasados, a vencer e de IR, e o total em aberto."""
    referencia = _data(data_referencia) or hoje()
    with _sessao() as sessao:
        contagem = servico.contar_pendentes(sessao, referencia)
        em_aberto = servico.consultar(sessao, servico.Filtro(data_referencia=referencia))
        return {
            "data_referencia": _plano(referencia),
            "atrasados": contagem.atrasados,
            "a_vencer": contagem.a_vencer,
            "imposto_de_renda": contagem.imposto_de_renda,
            "total_em_aberto": _plano(servico.total_de(em_aberto)),
            "total_em_aberto_formatado": formatar_moeda(servico.total_de(em_aberto)),
        }


@_ferramenta(LEITURA)
def buscar_clientes(
    termo: str = "", tipo_pessoa: str | None = None, limite: int = 50
) -> list[dict]:
    """Acha clientes por nome (sem acento), código ou CPF/CNPJ. tipo_pessoa: PF ou PJ."""
    with _sessao() as sessao:
        achados = clientes.buscar(sessao, termo, _enum(TipoPessoa, tipo_pessoa), limite=limite)
        return [_dados_cliente(c) for c in achados]


@_ferramenta(LEITURA)
def obter_cliente(cliente_id: int | None = None, codigo: str | None = None) -> dict:
    """Ficha do cliente (por id ou código): contatos, vínculo com empresa e pessoas vinculadas."""
    with _sessao() as sessao:
        c = _cliente(sessao, cliente_id) if cliente_id else (
            clientes.por_codigo(sessao, codigo) if codigo else None
        )
        if c is None:
            raise ValueError("Informe cliente_id ou um código que exista.")
        dados = _dados_cliente(c)
        dados["contatos"] = [
            {"id": x.id, "nome": x.nome, "tratamento": x.nome_tratamento,
             "telefone": x.telefone, "email": x.email, "principal": x.principal,
             "ativo": x.ativo}
            for x in c.contatos
        ]
        dados["empresa"] = c.empresa.nome if c.empresa else None
        dados["vinculados"] = [
            {"id": v.id, "nome": v.nome} for v in clientes.vinculados(sessao, c)
        ]
        return dados


@_ferramenta(LEITURA)
def listar_honorarios(
    situacao: str | None = "EM_ABERTO",
    termo: str = "",
    cliente_id: int | None = None,
    incluir_vinculados: bool = False,
    tipo_pessoa: str | None = None,
    vencimento_de: str | None = None,
    vencimento_ate: str | None = None,
    pago_de: str | None = None,
    pago_ate: str | None = None,
    atraso_minimo: int | None = None,
    atraso_maximo: int | None = None,
    data_referencia: str | None = None,
    limite: int = 100,
) -> dict:
    """Honorários com juros calculados. situacao: EM_ABERTO, PAGO, CANCELADO, RENEGOCIADO ou
    vazio para todas. Já vêm ordenados por vencimento; ``total`` soma tudo que casou, mesmo
    acima do ``limite``."""
    with _sessao() as sessao:
        filtro = servico.Filtro(
            termo=termo,
            cliente_id=cliente_id,
            incluir_vinculados_de=cliente_id if incluir_vinculados else None,
            tipo_pessoa=_enum(TipoPessoa, tipo_pessoa),
            status=_enum(StatusPagamento, situacao),
            vencimento_de=_data(vencimento_de),
            vencimento_ate=_data(vencimento_ate),
            pago_de=_data(pago_de),
            pago_ate=_data(pago_ate),
            atraso_minimo=atraso_minimo,
            atraso_maximo=atraso_maximo,
            data_referencia=_data(data_referencia),
        )
        linhas = servico.consultar(sessao, filtro)
        return {
            "quantidade": len(linhas),
            "total": _plano(servico.total_de(linhas)),
            "mostrando": min(len(linhas), limite),
            "honorarios": [_dados_linha(linha) for linha in linhas[:limite]],
        }


@_ferramenta(LEITURA)
def obter_honorario(honorario_id: int, data_referencia: str | None = None) -> dict:
    """Um honorário com pagamentos, cobranças e o cálculo de juros na data."""
    with _sessao() as sessao:
        h = _honorario(sessao, honorario_id)
        referencia = _data(data_referencia) or hoje()
        linha = servico.montar_linha(h, calculo.atualizar(h, referencia), referencia)
        dados = _dados_linha(linha)
        dados["conta_dos_juros"] = calculo.explicar(linha.atualizacao)
        dados["pagamentos"] = [_dados_pagamento(p) for p in h.pagamentos]
        dados["historico_cobrancas"] = [
            {"id": c.id, "data_hora": _plano(c.data_hora), "canal": c.canal.value,
             "resultado": c.resultado.value, "promessa": _plano(c.promessa_data),
             "observacao": c.observacao, "estornada": c.estornada_em is not None}
            for c in h.cobrancas
        ]
        return dados


@_ferramenta(LEITURA)
def historico_de_cobrancas(cliente_id: int) -> list[dict]:
    """Todas as cobranças feitas a um cliente, da mais recente para a mais antiga."""
    with _sessao() as sessao:
        return [
            {"id": c.id, "honorario_id": c.honorario_id, "data_hora": _plano(c.data_hora),
             "canal": c.canal.value, "resultado": c.resultado.value,
             "promessa": _plano(c.promessa_data), "observacao": c.observacao,
             "estornada": c.estornada_em is not None}
            for c in cobrancas.historico_do_cliente(sessao, cliente_id)
        ]


@_ferramenta(LEITURA)
def previa_de_pagamento(honorario_id: int, data: str | None = None) -> dict:
    """O que a baixa mostraria: saldo, juros devidos e total na data (hoje se vazia)."""
    with _sessao() as sessao:
        return _plano(pagamentos.previa(_honorario(sessao, honorario_id), _data(data)))


@_ferramenta(LEITURA)
def ranking_de_devedores(criterio: str = "maiores", limite: int = 10) -> dict:
    """Ranking de quem deve. criterio: maiores (valor), antigos (há mais tempo),
    quantidade (mais honorários em aberto) ou cobrados (mais cobranças)."""
    with _sessao() as sessao:
        if criterio == "cobrados":
            return _plano(rankings.mais_cobrados(sessao, limite))
        grupos = servico.agrupar_por_cliente(servico.consultar(sessao, servico.Filtro()))
        funcoes = {
            "maiores": lambda: rankings.maiores_devedores(grupos, limite),
            "antigos": lambda: rankings.devendo_ha_mais_tempo(grupos, limite, hoje()),
            "quantidade": lambda: rankings.mais_honorarios_em_aberto(grupos, limite),
        }
        if criterio not in funcoes:
            raise ValueError("criterio: maiores, antigos, quantidade ou cobrados.")
        return _plano(funcoes[criterio]())


@_ferramenta(LEITURA)
def listar_tipos_e_grupos(incluir_inativos: bool = False) -> dict:
    """Grupos e tipos de honorário (com ids), formas de envio e formas de pagamento válidas."""
    with _sessao() as sessao:
        return {
            "grupos": [{"id": g.id, "nome": g.nome, "ativo": g.ativo}
                       for g in tipos.listar_grupos(sessao, not incluir_inativos)],
            "tipos": [{"id": t.id, "nome": t.nome, "grupo_id": t.grupo_id, "ativo": t.ativo}
                      for t in tipos.listar_tipos(sessao, not incluir_inativos)],
            "formas_de_envio": [{"id": f.id, "nome": f.nome}
                                for f in tipos.listar_formas(sessao)],
            "formas_de_pagamento": [item.value for item in FormaPagamento],
        }


@_ferramenta(LEITURA)
def listar_contratos(cliente_id: int | None = None, incluir_suspensos: bool = True) -> dict:
    """Contratos de honorário recorrente, com o total mensal."""
    with _sessao() as sessao:
        lista = contratos.listar(sessao, cliente_id, incluir_suspensos)
        return {
            "total_mensal": _plano(contratos.total_mensal(sessao)),
            "contratos": [
                {"id": c.id, "cliente_id": c.cliente_id, "cliente": c.cliente.nome,
                 "tipo_id": c.tipo_id, "tipo": c.tipo.nome, "valor": _plano(c.valor),
                 "periodicidade": c.periodicidade.value, "dia_vencimento": c.dia_vencimento,
                 "suspenso": c.suspenso}
                for c in lista
            ],
        }


@_ferramenta(LEITURA)
def ler_configuracao(chave: str, padrao: str = "") -> str:
    """Lê uma configuração do sistema pela chave."""
    with _sessao() as sessao:
        return configuracoes.obter(sessao, chave, padrao)


# ---------------------------------------------------------------------- clientes
@_ferramenta(GRAVA)
def criar_cliente(
    codigo: str,
    nome: str,
    juros_dia: str,
    tipo_pessoa: str = "PJ",
    documento: str = "",
    observacoes: str = "",
    empresa_id: int | None = None,
) -> dict:
    """Cadastra um cliente. ``juros_dia`` é o juros de atraso em R$ por dia (obrigatório, sem
    padrão). Pessoa física pode ser vinculada (empresa_id) à empresa PJ que paga por ela."""
    with _sessao(gravar=True) as sessao:
        criado = clientes.criar(sessao, clientes.DadosCliente(
            codigo=codigo, nome=nome, documento=documento,
            tipo_pessoa=_enum(TipoPessoa, tipo_pessoa), observacoes=observacoes,
            juros_dia=_valor(juros_dia), empresa_id=empresa_id,
        ))
        return _dados_cliente(criado)


@_ferramenta(GRAVA)
def atualizar_cliente(
    cliente_id: int,
    nome: str | None = None,
    codigo: str | None = None,
    documento: str | None = None,
    situacao: str | None = None,
    subtipo: str | None = None,
    observacoes: str | None = None,
    empresa_id: int | None = None,
    desvincular_empresa: bool = False,
) -> dict:
    """Altera só os campos informados; os contatos ficam como estão. empresa_id vincula à
    empresa; desvincular_empresa tira o vínculo. Para o juros use alterar_juros_cliente."""
    with _sessao(gravar=True) as sessao:
        c = _cliente(sessao, cliente_id)
        dados = clientes.DadosCliente(
            codigo=codigo or c.codigo,
            nome=nome or c.nome,
            documento=c.documento if documento is None else documento,
            tipo_pessoa=c.tipo_pessoa,
            subtipo=_enum(SubtipoCliente, subtipo, c.subtipo),
            situacao=_enum(SituacaoCliente, situacao, c.situacao),
            observacoes=c.observacoes if observacoes is None else observacoes,
            # a camada de serviço trata a lista como a verdade: sem repassar os
            # contatos atuais, eles seriam apagados
            contatos=[
                clientes.DadosContato(
                    id=x.id, nome=x.nome, nome_tratamento=x.nome_tratamento,
                    telefone=x.telefone, email=x.email, principal=x.principal,
                    ativo=x.ativo, observacoes=x.observacoes,
                )
                for x in c.contatos
            ],
            empresa_id=None if desvincular_empresa else (empresa_id or c.empresa_id),
        )
        return _dados_cliente(clientes.atualizar(sessao, c, dados))


@_ferramenta(GRAVA)
def alterar_juros_cliente(
    cliente_id: int, juros_dia: str, vigencia_inicio: str | None = None, motivo: str = ""
) -> dict:
    """Novo juros de atraso (R$ por dia) do cliente, valendo da vigência em diante (hoje se
    vazia); o que já venceu antes continua com o juros da época."""
    with _sessao(gravar=True) as sessao:
        c = _cliente(sessao, cliente_id)
        clientes.registrar_juros(sessao, c, _valor(juros_dia), _data(vigencia_inicio), motivo)
        return _dados_cliente(c)


@_ferramenta(APAGA)
def excluir_cliente(cliente_id: int, motivo: str = "") -> str:
    """Exclusão lógica (o histórico fica). Recusa se houver honorário em aberto."""
    with _sessao(gravar=True) as sessao:
        c = _cliente(sessao, cliente_id)
        clientes.excluir(sessao, c, motivo)
        return f"{c.nome} excluído."


@_ferramenta(GRAVA)
def suspender_cobranca(cliente_id: int, motivo: str, ate: str | None = None) -> dict:
    """Tira o cliente das filas de cobrança (a dívida e os juros continuam). Sem ``ate``,
    vale até retomar."""
    with _sessao(gravar=True) as sessao:
        c = _cliente(sessao, cliente_id)
        clientes.suspender_cobranca(sessao, c, motivo, _data(ate))
        return _dados_cliente(c)


@_ferramenta(GRAVA)
def retomar_cobranca(cliente_id: int, motivo: str = "") -> dict:
    """Volta a cobrar um cliente que estava com a cobrança suspensa."""
    with _sessao(gravar=True) as sessao:
        c = _cliente(sessao, cliente_id)
        clientes.retomar_cobranca(sessao, c, motivo)
        return _dados_cliente(c)


# ------------------------------------------------------------------- honorários
def _competencia(ano: int | None, mes: int | None) -> Competencia:
    return Competencia(ano=ano, mes=mes)


@_ferramenta(GRAVA)
def criar_honorario(
    cliente_id: int,
    tipo_id: int,
    valor: str,
    vencimento: str,
    competencia_ano: int | None = None,
    competencia_mes: int | None = None,
    descricao: str = "",
    meio: str = "BOLETO",
    numero_boleto: str = "",
    observacoes: str = "",
) -> dict:
    """Lança um honorário avulso. A competência (ano/mês) depende do tipo; use
    listar_tipos_e_grupos para os ids."""
    with _sessao(gravar=True) as sessao:
        criado = servico.criar(sessao, servico.DadosHonorario(
            cliente_id=cliente_id, tipo_id=tipo_id, valor=_valor(valor),
            vencimento=_data(vencimento),
            competencia=_competencia(competencia_ano, competencia_mes),
            descricao=descricao, meio=_enum(MeioCobranca, meio), numero_boleto=numero_boleto,
            observacoes=observacoes,
        ))
        return _dados_linha(servico.montar_linha(criado, None, hoje()))


@_ferramenta(GRAVA)
def atualizar_honorario(
    honorario_id: int,
    valor: str | None = None,
    vencimento: str | None = None,
    descricao: str | None = None,
    numero_boleto: str | None = None,
    observacoes: str | None = None,
) -> dict:
    """Corrige valor, vencimento ou textos de um honorário em aberto (pago não edita)."""
    with _sessao(gravar=True) as sessao:
        h = _honorario(sessao, honorario_id)
        atualizado = servico.atualizar(sessao, h, servico.DadosHonorario(
            cliente_id=h.cliente_id, tipo_id=h.tipo_id,
            valor=_valor(valor) if valor else h.valor_original,
            vencimento=_data(vencimento) or h.vencimento,
            competencia=_competencia(h.competencia_ano, h.competencia_mes),
            descricao=h.descricao if descricao is None else descricao,
            meio=h.meio, forma_envio_id=h.forma_envio_id,
            observacoes=h.observacoes if observacoes is None else observacoes,
            numero_boleto=h.numero_boleto if numero_boleto is None else numero_boleto,
            contrato_id=h.contrato_id, situacao_envio=h.situacao_envio, enviado_em=h.enviado_em,
        ))
        return _dados_linha(servico.montar_linha(atualizado, None, hoje()))


@_ferramenta(APAGA)
def cancelar_honorario(honorario_id: int, motivo: str) -> str:
    """Cancela um honorário (zera o saldo). Pago precisa de estorno antes."""
    with _sessao(gravar=True) as sessao:
        servico.cancelar(sessao, _honorario(sessao, honorario_id), motivo)
        return f"Honorário {honorario_id} cancelado."


# -------------------------------------------------------------------- pagamentos
@_ferramenta(GRAVA)
def registrar_pagamento(
    honorario_id: int,
    valor_pago: str,
    data: str | None = None,
    juros_cobrados: str | None = None,
    desconto: str = "0",
    justificativa_desconto: str = "",
    forma_pagamento: str = "PIX",
    observacoes: str = "",
) -> dict:
    """Dá baixa (total ou parcial). Sem ``juros_cobrados`` cobra os juros devidos; desconto
    exige justificativa. Use previa_de_pagamento antes para ver os valores."""
    with _sessao(gravar=True) as sessao:
        h = _honorario(sessao, honorario_id)
        pago = pagamentos.registrar(sessao, h, pagamentos.DadosPagamento(
            data=_data(data) or hoje(), valor_pago=_valor(valor_pago),
            juros_cobrados=_valor(juros_cobrados) if juros_cobrados else None,
            desconto=_valor(desconto), justificativa_desconto=justificativa_desconto,
            forma_pagamento=_enum(FormaPagamento, forma_pagamento), observacoes=observacoes,
        ))
        return {"pagamento": _dados_pagamento(pago), "honorario_status": h.status.value}


@_ferramenta(APAGA)
def estornar_pagamento(pagamento_id: int, motivo: str) -> str:
    """Desfaz um pagamento (o registro fica guardado e o honorário volta para em aberto)."""
    with _sessao(gravar=True) as sessao:
        pago = sessao.get(Pagamento, pagamento_id)
        if pago is None:
            raise ValueError(f"Pagamento {pagamento_id} não encontrado.")
        pagamentos.estornar(sessao, pago, motivo)
        return f"Pagamento {pagamento_id} estornado."


# ---------------------------------------------------------------------- cobranças
@_ferramenta(GRAVA)
def registrar_cobranca(
    honorario_ids: list[int],
    canal: str = "WHATSAPP",
    resultado: str = "AGUARDANDO",
    observacao: str = "",
    data: str | None = None,
) -> str:
    """Registra que os honorários foram cobrados (a mensagem foi enviada por fora).
    resultado: AGUARDANDO, SEM_RESPOSTA, PROMETEU_PAGAR, PEDIU_RENEGOCIACAO, CONTESTOU,
    PAGOU ou OUTRO. Para PROMETEU_PAGAR use registrar_promessa."""
    resultado_enum = _enum(ResultadoCobranca, resultado)
    if resultado_enum is ResultadoCobranca.PROMETEU_PAGAR:
        raise ValueError("Para promessa de pagamento use registrar_promessa (precisa da data).")
    with _sessao(gravar=True) as sessao:
        referencia = hoje()
        linhas = [_linha(sessao, i, referencia) for i in honorario_ids]
        total = cobrancas.registrar_sem_mensagem(
            sessao, linhas, observacao, _enum(Canal, canal), resultado_enum, _data(data),
        )
        return f"{total} cobrança(s) registrada(s)."


@_ferramenta(GRAVA)
def registrar_promessa(cobranca_id: int, promessa_data: str, observacao: str | None = None) -> str:
    """Marca uma cobrança já registrada como 'prometeu pagar' em tal data."""
    with _sessao(gravar=True) as sessao:
        c = sessao.get(Cobranca, cobranca_id)
        if c is None:
            raise ValueError(f"Cobrança {cobranca_id} não encontrada.")
        cobrancas.atualizar_resultado(
            sessao, c, ResultadoCobranca.PROMETEU_PAGAR, _data(promessa_data), observacao
        )
        return f"Promessa para {promessa_data} registrada."


@_ferramenta(APAGA)
def estornar_cobranca(cobranca_id: int, motivo: str) -> str:
    """Marca uma cobrança como estornada (registrada por engano)."""
    with _sessao(gravar=True) as sessao:
        c = sessao.get(Cobranca, cobranca_id)
        if c is None:
            raise ValueError(f"Cobrança {cobranca_id} não encontrada.")
        cobrancas.estornar(sessao, c, motivo)
        return f"Cobrança {cobranca_id} estornada."


@_ferramenta(LEITURA)
def gerar_mensagem_de_cobranca(cliente_id: int, honorario_ids: list[int] | None = None) -> str:
    """Texto pronto de cobrança do cliente (todos os honorários cobráveis, ou só os
    informados). Não envia nem registra nada."""
    with _sessao() as sessao:
        c = _cliente(sessao, cliente_id)
        referencia = hoje()
        alvo = cobrancas.linhas_em_aberto_do_cliente(sessao, cliente_id, referencia) \
            if not honorario_ids else [_linha(sessao, i, referencia) for i in honorario_ids]
        if not alvo:
            raise ValueError("Esse cliente não tem honorário em aberto para cobrar.")
        return cobrancas.gerar_texto(sessao, cobrancas.Selecao(cliente=c, linhas=list(alvo)))


# ---------------------------------------------------------------------- contratos
@_ferramenta(GRAVA)
def criar_contrato(
    cliente_id: int,
    tipo_id: int,
    valor: str,
    periodicidade: str = "MENSAL",
    dia_vencimento: int = 10,
    vigencia_inicio: str | None = None,
    vigencia_fim: str | None = None,
    observacoes: str = "",
) -> dict:
    """Cria o contrato recorrente de um cliente para um tipo de honorário."""
    with _sessao(gravar=True) as sessao:
        c = contratos.criar(sessao, contratos.DadosContrato(
            cliente_id=cliente_id, tipo_id=tipo_id, valor=_valor(valor),
            periodicidade=_enum(Periodicidade, periodicidade), dia_vencimento=dia_vencimento,
            vigencia_inicio=_data(vigencia_inicio), vigencia_fim=_data(vigencia_fim),
            observacoes=observacoes,
        ))
        return {"id": c.id, "valor": _plano(c.valor)}


@_ferramenta(GRAVA)
def reajustar_contrato(contrato_id: int, novo_valor: str, motivo: str = "") -> dict:
    """Novo valor do contrato daqui para frente (o que já foi gerado não muda)."""
    with _sessao(gravar=True) as sessao:
        contrato = sessao.get(contratos.Contrato, contrato_id)
        if contrato is None:
            raise ValueError(f"Contrato {contrato_id} não encontrado.")
        contratos.reajustar(sessao, contrato, _valor(novo_valor), motivo=motivo)
        return {"id": contrato.id, "valor": _plano(contrato.valor)}


@_ferramenta(GRAVA)
def suspender_contrato(contrato_id: int, suspenso: bool = True) -> str:
    """Pausa (ou retoma, com suspenso=false) a geração mensal de um contrato."""
    with _sessao(gravar=True) as sessao:
        contrato = sessao.get(contratos.Contrato, contrato_id)
        if contrato is None:
            raise ValueError(f"Contrato {contrato_id} não encontrado.")
        contratos.suspender(sessao, contrato, suspenso)
        return "Contrato suspenso." if suspenso else "Contrato retomado."


@_ferramenta(APAGA)
def excluir_contrato(contrato_id: int) -> str:
    """Exclui o contrato (lógico). Os honorários já gerados continuam."""
    with _sessao(gravar=True) as sessao:
        contrato = sessao.get(contratos.Contrato, contrato_id)
        if contrato is None:
            raise ValueError(f"Contrato {contrato_id} não encontrado.")
        contratos.excluir(sessao, contrato)
        return f"Contrato {contrato_id} excluído."


# ------------------------------------------------------------------ configuração
@_ferramenta(GRAVA)
def gravar_configuracao(chave: str, valor: str) -> str:
    """Grava uma configuração do sistema (chave/valor em texto)."""
    with _sessao(gravar=True) as sessao:
        configuracoes.gravar(sessao, chave, valor)
        return f"{chave} gravada."


# -------------------------------------------------------- e-mail do financeiro
def _email_guardado(sessao: Session, email_id: int):
    from .infra.modelos import EmailMensagem

    mensagem = sessao.get(EmailMensagem, email_id)
    if mensagem is None:
        raise ValueError(f"E-mail {email_id} não encontrado (use listar_emails).")
    return mensagem


def _conferencia(sessao: Session, mensagem, referencia: date):
    from .servicos import caixa_email
    from .servicos import email_cobrancas as servico_email

    leitura = servico_email.ler(caixa_email.recebido_de(mensagem).texto)
    if not (leitura.linhas or leitura.incompletas):
        return None
    return servico_email.conferir(sessao, leitura, referencia, caixa_email.correcoes(mensagem))


def _linha_lida(lida) -> dict:
    return {"cliente": lida.cliente, "codigo": lida.codigo or None,
            "competencia": lida.competencia, "vencimento": _plano(lida.vencimento),
            "valor": _plano(lida.valor), "notas": lida.notas}


@_ferramenta(GRAVA)
def verificar_emails() -> dict:
    """Busca na caixa (Outlook/Gmail, pasta "Cobranças") os e-mails novos da conversa com o
    financeiro e guarda no sistema, como o botão "Verificar caixa"."""
    from .servicos import caixa_email
    from .servicos import email_cobrancas as servico_email

    with _sessao() as sessao:
        caixa = caixa_email.abrir_caixa(sessao)
        conta = caixa_email.conta(sessao).provedor
        ja_tenho = caixa_email.ja_baixados(sessao, conta)
        assunto = servico_email.remetente(sessao).assunto
    try:   # a rede fica fora da sessão do banco
        resultado = caixa_email.verificar(caixa, conta, ja_tenho, assunto)
    except ErroDeNegocio as falha:
        raise ValueError(str(falha)) from falha
    with _sessao(gravar=True) as sessao:
        novas = caixa_email.guardar(sessao, resultado.conta, resultado.mensagens)
        caixa_email.registrar_verificacao(sessao)
        return {"novos": len(novas),
                "emails": [{"id": m.id, "assunto": m.assunto, "de": m.de,
                            "recebido_em": _plano(m.recebido_em)} for m in novas]}


@_ferramenta(LEITURA)
def listar_emails(limite: int = 20, so_nao_respondidos: bool = False) -> list[dict]:
    """Os e-mails da conversa com o financeiro guardados no sistema, do mais novo ao mais
    antigo (recebidos e enviados daqui)."""
    from .servicos import caixa_email

    with _sessao() as sessao:
        return [
            {"id": m.id, "assunto": m.assunto, "de": m.de, "para": m.para,
             "recebido_em": _plano(m.recebido_em), "enviado_por_voce": m.enviada,
             "lido": m.lida_em is not None, "respondido_em": _plano(m.respondida_em),
             "origem": m.conta}
            for m in caixa_email.listar(sessao, limite * 3)
            if not (so_nao_respondidos and (m.enviada or m.respondida_em))
        ][:limite]


@_ferramenta(GRAVA)
def ler_email(email_id: int, data_referencia: str | None = None) -> dict:
    """Abre um e-mail do financeiro (marca como lido) e já confere a lista dele com o
    sistema: quem continua em aberto, quem saiu da lista (provavelmente pagou: candidato a
    baixa), quem já está pago e as linhas que não bateram com nenhum honorário."""
    from .servicos import caixa_email

    with _sessao(gravar=True) as sessao:
        mensagem = _email_guardado(sessao, email_id)
        caixa_email.marcar_lido(sessao, mensagem)
        dados = {
            "id": mensagem.id, "assunto": mensagem.assunto, "de": mensagem.de,
            "para": mensagem.para, "cc": mensagem.cc,
            "recebido_em": _plano(mensagem.recebido_em),
            "texto": caixa_email.recebido_de(mensagem).texto,
            "respondido_em": _plano(mensagem.respondida_em),
        }
        conferencia = _conferencia(sessao, mensagem, _data(data_referencia) or hoje())
        if conferencia is None:
            dados["conferencia"] = None
            return dados
        dados["conferencia"] = {
            "data_do_email": _plano(conferencia.data),
            "continuam": [dict(_linha_lida(c.lida), honorario_id=c.linha.honorario.id)
                          for c in conferencia.continuam],
            "sairam_da_lista": [_dados_linha(linha) for linha in conferencia.sairam],
            "ja_pagos_no_sistema": [dict(_linha_lida(c.lida), honorario_id=c.linha.honorario.id)
                                    for c in conferencia.pagos],
            "nao_encontradas": [_linha_lida(lida) for lida in conferencia.nao_encontradas],
        }
        return dados


@_ferramenta(GRAVA)
def baixar_quem_saiu_da_lista(
    email_id: int,
    honorario_ids: list[int],
    data_pagamento: str | None = None,
    com_juros: bool = True,
    forma_pagamento: str = "BOLETO",
) -> dict:
    """Dá baixa nos honorários que saíram da lista do financeiro (use os ids de
    ``sairam_da_lista`` de ler_email). Com juros: valor atualizado até a data do pagamento;
    sem juros: só o saldo, e o juros vira abono justificado pela lista. Tudo ou nada."""
    from .servicos import email_cobrancas as servico_email

    with _sessao(gravar=True) as sessao:
        mensagem = _email_guardado(sessao, email_id)
        conferencia = _conferencia(sessao, mensagem, hoje())
        data_do_email = conferencia.data if conferencia else hoje()
        honorarios = [_honorario(sessao, i) for i in honorario_ids]
        quando = _data(data_pagamento) or data_do_email
        itens = servico_email.previa_baixa(honorarios, quando, com_juros, data_do_email)
        quantos = servico_email.dar_baixa(
            sessao, itens, quando, _enum(FormaPagamento, forma_pagamento), data_do_email,
        )
        return {"baixados": quantos, "itens": [
            {"honorario_id": item.honorario.id, "cliente": item.cliente.nome,
             "valor": _plano(item.valor), "juros_cobrados": _plano(item.juros_cobrados)}
            for item in itens
        ]}


def _email_de_resposta(sessao: Session, email_id: int | None, referencia: date):
    """A resposta como a tela monta: à lista do e-mail, ou nova com os vencidos do sistema."""
    from .servicos import caixa_email
    from .servicos import email_cobrancas as servico_email

    if email_id is None:
        email = servico_email.montar(sessao, referencia)
        email.assunto = servico_email.remetente(sessao).assunto   # e-mail novo, sem "RE:"
        if not email.para:
            email.para = caixa_email.ultimo_remetente(sessao)
        return email, None
    mensagem = _email_guardado(sessao, email_id)
    conferencia = _conferencia(sessao, mensagem, referencia)
    if conferencia is None:
        raise ValueError("Esse e-mail não tem a lista de honorários para responder.")
    email = servico_email.responder(sessao, conferencia, referencia,
                                    caixa_email.recebido_de(mensagem))
    return email, mensagem


def _com_introducao(email, introducao: str) -> tuple[str, str]:
    from html import escape

    from .servicos import email_cobrancas as servico_email

    html, texto = servico_email.html(email), servico_email.texto(email)
    if introducao.strip():
        paragrafos = "".join(f"<p>{escape(p)}</p>" for p in introducao.strip().split("\n\n"))
        html = paragrafos + html
        texto = introducao.strip() + "\n\n" + texto
    return html, texto


@_ferramenta(LEITURA)
def montar_resposta_financeiro(
    email_id: int | None = None, introducao: str = "", data_referencia: str | None = None
) -> dict:
    """Monta (sem enviar) a resposta ao financeiro: a lista dele sem quem saiu, com o
    andamento das cobranças por cliente. Sem email_id, um e-mail novo com os vencidos do
    sistema. Devolve para/cc/assunto, o corpo em HTML e texto e, quando o e-mail veio do
    Outlook, o id dele (outlook_message_id) para responder ou criar rascunho por outro meio."""
    with _sessao() as sessao:
        email, mensagem = _email_de_resposta(sessao, email_id, _data(data_referencia) or hoje())
        html, texto = _com_introducao(email, introducao)
        return {
            "para": email.para, "cc": email.cc, "assunto": email.assunto,
            "em_resposta_a": email.em_resposta_a or None,
            "outlook_message_id": (mensagem.ref_provedor or None) if mensagem else None,
            "html": html, "texto": texto,
        }


@_ferramenta(GRAVA)
def responder_financeiro(
    email_id: int | None = None,
    modo: str = "rascunho",
    introducao: str = "",
    para: str | None = None,
    cc: str | None = None,
    assunto: str | None = None,
    com_assinatura: bool = True,
    abrir_rascunho: bool = True,
) -> dict:
    """Responde ao financeiro com a lista atualizada (a mesma de montar_resposta_financeiro).

    modo="enviar": envia pela conta configurada no sistema (Outlook/Gmail), citando o
    e-mail dele, e guarda nos enviados — não dá para desfazer.
    modo="rascunho": grava um .eml ligado à conversa e (com abrir_rascunho) abre no
    Outlook, pronto para revisar e enviar. Nada sai sem você clicar em Enviar.
    """
    from .config import caminhos
    from .servicos import caixa_email
    from .servicos import email_cobrancas as servico_email

    if modo not in ("enviar", "rascunho"):
        raise ValueError('modo: "enviar" ou "rascunho".')
    with _sessao(gravar=True) as sessao:
        email, mensagem = _email_de_resposta(sessao, email_id, hoje())
        email.para = email.para if para is None else para
        email.cc = email.cc if cc is None else cc
        email.assunto = email.assunto if assunto is None else assunto
        html, texto = _com_introducao(email, introducao)

        if modo == "rascunho":
            bruto = servico_email.eml(email)
            if introducao.strip():   # o .eml é montado do e-mail; a introdução vai junto
                from email import policy
                from email.parser import BytesParser

                rascunho = BytesParser(policy=policy.default).parsebytes(bruto)
                rascunho.get_body(("plain",)).set_content(texto)
                rascunho.get_body(("html",)).set_content(html, subtype="html")
                bruto = bytes(rascunho)
            destino = caminhos.pasta_exportacoes() / \
                f"resposta-cobrancas-{email.referencia:%Y-%m-%d}.eml"
            destino.parent.mkdir(parents=True, exist_ok=True)
            destino.write_bytes(bruto)
            if abrir_rascunho:
                _abrir(str(destino))
            return {"rascunho": str(destino), "para": email.para, "assunto": email.assunto}

        if mensagem is not None and mensagem.conta not in (
            caixa_email.GMAIL, caixa_email.IMAP, caixa_email.MICROSOFT,
        ):
            raise ValueError("Esse e-mail veio de arquivo: use modo=\"rascunho\".")
        caixa = caixa_email.abrir_caixa(sessao)
        resposta = caixa_email.Resposta(
            para=email.para, cc=email.cc, assunto=email.assunto, html=html, texto=texto,
            assinatura=caixa_email.assinatura(sessao) if com_assinatura else b"",
        )
        if mensagem is None:
            enviada = caixa.enviar_novo(resposta)
            caixa_email.registrar_nova_enviada(sessao, caixa_email.conta(sessao).provedor,
                                               enviada)
        else:
            enviada = caixa.responder(caixa_email.original_de(mensagem), resposta)
            caixa_email.registrar_enviada(sessao, mensagem, enviada)
        return {"enviado": True, "para": email.para, "cc": email.cc, "assunto": email.assunto}


# --------------------------------------------------------------------- WhatsApp
def _abrir(alvo: str) -> None:
    """Abre link ou arquivo no computador (WhatsApp, Outlook), como um clique."""
    if sys.platform == "win32":
        import os

        os.startfile(alvo)   # noqa: S606 - é o comportamento pedido: abrir no Windows
        return
    import webbrowser

    if not webbrowser.open(alvo):
        raise ValueError(f"Não consegui abrir {alvo}.")


def _app_grafico():
    """Qt sem janela, para desenhar a imagem de cobrança e usar a área de transferência."""
    import os

    from PySide6.QtGui import QGuiApplication

    if QGuiApplication.instance() is None:
        if sys.platform.startswith("linux") and not os.environ.get("DISPLAY") \
                and not os.environ.get("WAYLAND_DISPLAY"):
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        _app_grafico.app = QGuiApplication(["financeiro-mcp"])
    return QGuiApplication.instance()


def _imagem_de_cobranca(sessao: Session, selecao, nome: str, vencimento: date | None = None):
    """Gera o PNG da imagem de cobrança (o mesmo da tela) na pasta de exportações."""
    from .config import caminhos
    from .servicos import imagem_cobranca
    from .ui.comum.imagem_cobranca import renderizar

    _app_grafico()
    dados = imagem_cobranca.montar(sessao, selecao, vencimento=vencimento)
    texto_svg = imagem_cobranca.caminho_fundo(sessao).read_text(encoding="utf-8")
    imagem = renderizar(texto_svg, dados)
    destino = caminhos.pasta_exportacoes() / f"{nome}.png"
    destino.parent.mkdir(parents=True, exist_ok=True)
    if not imagem.save(str(destino), "PNG"):
        raise ValueError(f"Não consegui salvar a imagem em {destino}.")
    return imagem, destino


def _copiar_imagem(imagem, arquivo) -> None:
    """Põe a imagem na área de transferência, para colar no WhatsApp com Ctrl+V.

    No Windows vai por um PowerShell à parte: a área de transferência do Qt depende do
    laço de mensagens, que o servidor (sem janela) não tem; o .NET copia e deixa a
    imagem lá mesmo depois de terminar.
    """
    if sys.platform != "win32":
        _app_grafico().clipboard().setImage(imagem)
        return
    import subprocess

    caminho = str(arquivo).replace("'", "''")
    script = ("Add-Type -AssemblyName System.Windows.Forms, System.Drawing; "
              f"$img = [System.Drawing.Image]::FromFile('{caminho}'); "
              "[System.Windows.Forms.Clipboard]::SetImage($img); $img.Dispose()")
    subprocess.run(
        ["powershell", "-NoProfile", "-STA", "-Command", script],
        check=True, timeout=30, capture_output=True,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )


def _rotulo_linha(linha) -> str:
    h = linha.honorario
    return f"{h.id} · {h.tipo.nome} {h.competencia_rotulo or h.descricao}".strip()


@_ferramenta(GRAVA)
def cobrar_por_whatsapp(
    cliente_id: int,
    honorario_ids: list[int] | None = None,
    contato_id: int | None = None,
    texto: str | None = None,
    chave_pix_no_fim: bool = True,
    com_imagem: bool = True,
    abrir: bool = True,
    registrar: bool = False,
    atualizado_ate: str | None = None,
    vencimento: str | None = None,
    juros_manuais: dict[str, str] | None = None,
    lembrete: bool = False,
) -> dict:
    """Prepara a cobrança no WhatsApp do cliente (telefone do contato principal, ou do
    contato_id): a mensagem do modelo (ou ``texto``) com a chave Pix sozinha na última
    linha — para o cliente copiar, já que da imagem não dá — e abre o WhatsApp com tudo
    escrito. Com imagem, a imagem de cobrança (com o QR do Pix) fica copiada: no WhatsApp,
    Ctrl+V cola a imagem. O WhatsApp não deixa enviar sozinho: quem aperta Enviar é você.
    ``registrar`` grava a cobrança (só quando ela for de fato enviada).

    Como na tela Cobrar: ``atualizado_ate`` é a data até onde vão os juros (hoje se
    vazia); ``vencimento`` é o "pagar até" da imagem e do lembrete — não mexe nos juros;
    ``juros_manuais`` ({id do honorário: juros}) cobra outro juros no lugar do
    calculado; ``lembrete`` usa o modelo de lembrete (ainda no prazo)."""
    from .servicos import configuracoes as servico_config

    with _sessao(gravar=registrar) as sessao:
        c = _cliente(sessao, cliente_id)
        referencia = _data(atualizado_ate) or hoje()
        linhas = [_linha(sessao, i, referencia) for i in honorario_ids] if honorario_ids \
            else cobrancas.linhas_em_aberto_do_cliente(sessao, cliente_id, referencia)
        linhas = [linha for linha in linhas if linha.honorario.status.cobravel]
        if not linhas:
            raise ValueError(f"{c.nome} não tem honorário em aberto para cobrar.")
        for chave, juros in (juros_manuais or {}).items():
            linha = next((x for x in linhas if x.honorario.id == int(chave)), None)
            if linha is None:
                raise ValueError(f"O honorário {chave} não está nesta cobrança.")
            linha.atualizacao = calculo.com_juros_manual(linha.atualizacao, _valor(juros))
        contato = next((x for x in c.contatos if x.id == contato_id), None) if contato_id \
            else c.contato_principal
        if contato_id and contato is None:
            raise ValueError(f"O contato {contato_id} não é de {c.nome}.")
        prazo = _data(vencimento)
        selecao = cobrancas.Selecao(cliente=c, linhas=list(linhas), contato=contato,
                                    data_referencia=referencia, forcar_lembrete=lembrete,
                                    vencimento_lembrete=prazo)
        mensagem = (texto or cobrancas.gerar_texto(sessao, selecao)).strip()

        recebedor = servico_config.recebedor(sessao)
        if chave_pix_no_fim and recebedor.chave:
            ultima = mensagem.splitlines()[-1].strip() if mensagem else ""
            if ultima != recebedor.chave.strip():
                mensagem += f"\n\n{recebedor.chave.strip()}"

        link = cobrancas.link_whatsapp(sessao, contato, mensagem)
        resultado = {"cliente": c.nome, "telefone": contato.telefone if contato else None,
                     "contato": contato.tratamento_ou_nome if contato else None,
                     "mensagem": mensagem, "link": link, "imagem": None,
                     "total": _plano(selecao.total),
                     "conta_dos_juros": {_rotulo_linha(x): calculo.explicar(x.atualizacao)
                                         for x in linhas}}
        if com_imagem:
            imagem, destino = _imagem_de_cobranca(
                sessao, selecao, f"cobranca_{c.codigo}_{referencia:%Y-%m-%d}", prazo)
            try:
                _copiar_imagem(imagem, destino)
                resultado["imagem_copiada"] = True
            except Exception as falha:  # noqa: BLE001 - a imagem fica salva do mesmo jeito
                resultado["imagem_copiada"] = False
                resultado["aviso"] = f"Não consegui copiar a imagem ({falha}); ela está salva."
            resultado["imagem"] = str(destino)
        if abrir:
            _abrir(link)
        if registrar:
            cobrancas.registrar(sessao, selecao, texto=mensagem, canal=Canal.WHATSAPP)
            resultado["cobranca_registrada"] = True
        return resultado


# ---------------------------------------------------------------------- contatos
def _contato(sessao: Session, contato_id: int):
    from .infra.modelos import ClienteContato

    contato = sessao.get(ClienteContato, contato_id)
    if contato is None:
        raise ValueError(f"Contato {contato_id} não encontrado (veja obter_cliente).")
    return contato


def _dados_contato(contato) -> dict:
    return {"id": contato.id, "cliente_id": contato.cliente_id, "nome": contato.nome,
            "tratamento": contato.nome_tratamento, "telefone": contato.telefone,
            "email": contato.email, "principal": contato.principal, "ativo": contato.ativo,
            "observacoes": contato.observacoes}


@_ferramenta(GRAVA)
def adicionar_contato(
    cliente_id: int,
    nome: str = "",
    tratamento: str = "",
    telefone: str = "",
    email: str = "",
    observacoes: str = "",
    principal: bool = False,
) -> dict:
    """Cadastra um contato (responsável) no cliente. ``tratamento`` é como chamar na
    saudação da mensagem; ``principal`` faz dele o contato da cobrança."""
    from .servicos import contatos as servico_contatos

    with _sessao(gravar=True) as sessao:
        contato = servico_contatos.adicionar(
            sessao, _cliente(sessao, cliente_id), nome, tratamento, telefone, email,
            observacoes, principal)
        return _dados_contato(contato)


@_ferramenta(GRAVA)
def atualizar_contato(
    contato_id: int,
    nome: str | None = None,
    tratamento: str | None = None,
    telefone: str | None = None,
    email: str | None = None,
    observacoes: str | None = None,
    principal: bool | None = None,
    ativo: bool | None = None,
) -> dict:
    """Muda só os campos informados de um contato (ids em obter_cliente)."""
    from .servicos import contatos as servico_contatos

    with _sessao(gravar=True) as sessao:
        contato = _contato(sessao, contato_id)
        mudou = servico_contatos.atualizar(sessao, contato, nome, tratamento, telefone, email,
                                           observacoes, principal, ativo)
        return {**_dados_contato(contato), "mudou": mudou}


@_ferramenta(APAGA)
def remover_contato(contato_id: int) -> str:
    """Apaga o contato do cliente (outro ativo vira o principal, se preciso)."""
    from .servicos import contatos as servico_contatos

    with _sessao(gravar=True) as sessao:
        contato = _contato(sessao, contato_id)
        nome = contato.nome or contato.telefone
        servico_contatos.remover(sessao, contato)
        return f"Contato {nome} removido."


@_ferramenta(GRAVA)
def cadastrar_contatos_em_lote(contatos: list[dict], gravar: bool = False) -> dict:
    """Cruza uma lista de contatos das empresas com o cadastro e grava os que baterem.

    Cada item diz de qual empresa é por um destes (na ordem em que são tentados):
    ``codigo`` (código do cliente no sistema), ``documento`` (CPF/CNPJ),
    ``telefone_empresa`` (um telefone que a empresa já tem cadastrado) ou
    ``nome_empresa`` (nome exato). E traz o responsável: ``nome``, ``tratamento``,
    ``telefone``, ``email``, ``observacoes`` e ``principal`` (true/false).

    Telefone igual (com ou sem 55, com ou sem o 9) é o mesmo contato: é atualizado só no
    que vier preenchido; senão entra como novo. O que não achar a empresa, ou achar mais
    de uma, não grava e volta com o motivo.

    Com ``gravar=false`` (padrão) só mostra o que faria; confira e chame de novo com
    ``gravar=true`` — o cruzamento é refeito na hora de gravar.
    """
    from .servicos import contatos as servico_contatos

    campos = {f.name for f in fields(servico_contatos.ItemContato)}
    itens = []
    for numero, bruto in enumerate(contatos, start=1):
        if not isinstance(bruto, dict):
            raise ValueError(f"Item {numero} da lista não é um objeto.")
        desconhecidos = set(bruto) - campos
        if desconhecidos:
            raise ValueError(f"Item {numero}: campo(s) que não existem: "
                             f"{', '.join(sorted(desconhecidos))}. "
                             f"Use: {', '.join(sorted(campos))}.")
        dados = {chave: ("" if valor is None else valor) for chave, valor in bruto.items()}
        dados = {chave: (bool(valor) if chave == "principal" else str(valor).strip())
                 for chave, valor in dados.items()}
        itens.append(servico_contatos.ItemContato(**dados))

    with _sessao(gravar=gravar) as sessao:
        cruzamentos = servico_contatos.cruzar(sessao, itens)
        gravados = servico_contatos.gravar(sessao, cruzamentos) if gravar else 0
        contagem: dict[str, int] = {}
        for cruzamento in cruzamentos:
            contagem[cruzamento.acao] = contagem.get(cruzamento.acao, 0) + 1
        return {
            "gravado": gravar,
            "gravados": gravados,
            "resumo": contagem,
            "itens": [
                {"item": c.indice,
                 "cliente": f"{c.cliente.codigo} – {c.cliente.nome}" if c.cliente else None,
                 "cliente_id": c.cliente.id if c.cliente else None,
                 "achado_por": c.achado_por or None, "acao": c.acao,
                 "contato_existente_id": c.contato.id if c.contato else None,
                 "mudancas": c.mudancas, "motivo": c.motivo or None,
                 "responsavel": c.item.nome, "telefone": c.item.telefone}
                for c in cruzamentos
            ],
        }


# ------------------------------------------------------------ planilhas e layouts
def _caminho(arquivo: str):
    from pathlib import Path

    caminho = Path(arquivo).expanduser()
    if not caminho.is_file():
        raise ValueError(f"Arquivo não encontrado: {caminho}")
    return caminho


@_ferramenta(LEITURA)
def ler_planilha(
    arquivo: str, aba: str | None = None, linha_inicial: int = 1, limite: int = 500
) -> dict:
    """Lê uma planilha do computador (.xlsx, .xlsm ou .csv) e devolve as linhas com a
    letra de cada coluna, para entender uma lista antes de importar ou cruzar."""
    caminho = _caminho(arquivo)
    linhas: list[dict] = []
    if caminho.suffix.lower() == ".csv":
        import csv

        from openpyxl.utils import get_column_letter

        texto = caminho.read_text(encoding="utf-8-sig", errors="replace")
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t") if texto else csv.excel
        leitor = csv.reader(texto.splitlines(), dialeto)
        abas, nome_aba = ["(csv)"], "(csv)"
        for numero, celulas in enumerate(leitor, start=1):
            if numero < linha_inicial or not any(c.strip() for c in celulas):
                continue
            linhas.append({"linha": numero, **{get_column_letter(i): c
                                               for i, c in enumerate(celulas, start=1) if c}})
            if len(linhas) > limite:
                break
    else:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        livro = load_workbook(caminho, data_only=True, read_only=True)
        try:
            abas = list(livro.sheetnames)
            nome_aba = aba or abas[0]
            if nome_aba not in abas:
                raise ValueError(f"A planilha não tem a aba {aba!r}. Abas: {', '.join(abas)}")
            for numero, celulas in enumerate(
                    livro[nome_aba].iter_rows(min_row=linha_inicial, values_only=True),
                    start=linha_inicial):
                valores = {get_column_letter(i): _plano(c)
                           for i, c in enumerate(celulas, start=1) if c not in (None, "")}
                if valores:
                    linhas.append({"linha": numero, **valores})
                if len(linhas) > limite:
                    break
        finally:
            livro.close()
    return {"abas": abas, "aba": nome_aba, "linhas": linhas[:limite],
            "mais_linhas": len(linhas) > limite}


@_ferramenta(LEITURA)
def campos_de_importacao() -> list[dict]:
    """As informações que um layout de importação pode ler, com a chave de cada uma."""
    from .servicos.importacao import personalizada

    return [{"chave": c.chave, "rotulo": c.rotulo, "grupo": c.grupo, "tipo": c.tipo,
             "dica": c.dica} for c in personalizada.CAMPOS]


@_ferramenta(LEITURA)
def listar_layouts_de_importacao() -> list[dict]:
    """Os layouts de importação salvos (os mesmos da tela Importação personalizada)."""
    from .servicos.importacao import personalizada

    with _sessao() as sessao:
        return [layout.para_dict() for layout in personalizada.listar_layouts(sessao).values()]


def _layout(campos: dict, nome: str = "", aba: str = "", linha_inicial: int = 2):
    from .servicos.importacao import personalizada

    mapas = {}
    for chave, valor in campos.items():
        if chave not in personalizada.POR_CHAVE:
            raise ValueError(f"Campo {chave!r} não existe (veja campos_de_importacao).")
        if isinstance(valor, str):
            valor = {"coluna": valor}
        if not isinstance(valor, dict):
            raise ValueError(f"{chave}: use a letra da coluna ou {{coluna, padrao}}.")
        mapas[chave] = personalizada.Mapeamento(
            coluna=str(valor.get("coluna") or "").strip().upper(),
            padrao=str(valor.get("padrao") or "").strip())
    return personalizada.Layout(nome=nome.strip(), aba=aba.strip(),
                                linha_inicial=linha_inicial, campos=mapas)


@_ferramenta(GRAVA)
def salvar_layout_de_importacao(
    nome: str, campos: dict, aba: str = "", linha_inicial: int = 2
) -> dict:
    """Cria (ou substitui) um layout de importação. ``campos`` liga cada chave de
    campos_de_importacao à letra da coluna — {"cliente_codigo": "A"} — ou a
    {"coluna": "B", "padrao": "valor fixo quando a célula vier vazia"}. Precisa da coluna
    do código, do CPF/CNPJ ou do nome do cliente."""
    from .servicos.importacao import personalizada

    with _sessao(gravar=True) as sessao:
        layout = _layout(campos, nome, aba, linha_inicial)
        personalizada.salvar_layout(sessao, layout)
        return personalizada.listar_layouts(sessao)[layout.nome].para_dict()


@_ferramenta(APAGA)
def excluir_layout_de_importacao(nome: str) -> str:
    """Apaga um layout de importação salvo."""
    from .servicos.importacao import personalizada

    with _sessao(gravar=True) as sessao:
        personalizada.excluir_layout(sessao, nome)
        return f"Layout {nome!r} excluído."


def _layout_salvo(sessao: Session, nome: str):
    from .servicos.importacao import personalizada

    layout = personalizada.listar_layouts(sessao).get(nome)
    if layout is None:
        raise ValueError(f"Layout {nome!r} não existe (veja listar_layouts_de_importacao).")
    return layout


@_ferramenta(GRAVA)
def exportar_modelo_de_layout(nome: str, destino: str | None = None) -> str:
    """Gera a planilha-modelo do layout (cabeçalhos nas colunas certas e instruções),
    para preencher e importar. Sem destino, vai para a pasta de exportações."""
    from .config import caminhos
    from .servicos.importacao import personalizada

    with _sessao() as sessao:
        layout = _layout_salvo(sessao, nome)
    alvo = destino or str(caminhos.pasta_exportacoes() / f"modelo-{nome}.xlsx")
    return str(personalizada.exportar_modelo(layout, alvo))


def _problema(p) -> dict:
    return {"linha": p.linha, "campo": p.campo, "valor": p.valor, "motivo": p.motivo,
            "gravidade": p.gravidade}


@_ferramenta(LEITURA)
def previa_de_importacao(
    arquivo: str, layout: str, atualizar_existentes: bool = True, limite: int = 200
) -> dict:
    """Lê a planilha com um layout salvo e confere tudo com o banco, sem gravar: o que
    cada linha faria com o cliente, o contrato e o honorário, e os problemas."""
    from .servicos.importacao import personalizada

    with _sessao() as sessao:
        previa = personalizada.ler(sessao, _caminho(arquivo), _layout_salvo(sessao, layout),
                                   atualizar_existentes)
        return {
            "resumo": previa.resumo(),
            "aba": previa.aba,
            "problemas_gerais": [_problema(p) for p in previa.problemas],
            "linhas": [
                {"linha": linha.linha, "cliente": linha.rotulo_cliente,
                 "cliente_acao": linha.acao_cliente or None,
                 "contrato_acao": linha.acao_contrato or None,
                 "honorario_acao": linha.acao_honorario or None,
                 "bloqueada": linha.bloqueada,
                 "problemas": [_problema(p) for p in linha.problemas]}
                for linha in previa.linhas[:limite]
            ],
            "mais_linhas": len(previa.linhas) > limite,
        }


@_ferramenta(GRAVA)
def importar_planilha(
    arquivo: str,
    layout: str,
    atualizar_existentes: bool = True,
    pular_linhas: list[int] | None = None,
) -> dict:
    """Importa a planilha com um layout salvo, como o botão Importar da tela: cria e
    atualiza clientes, contatos, contratos e honorários. Linhas com erro ficam de fora;
    ``pular_linhas`` tira outras (números das linhas do Excel). Rode previa_de_importacao
    antes e confirme com o usuário."""
    from .servicos.importacao import personalizada

    pular = set(pular_linhas or [])
    with _sessao(gravar=True) as sessao:
        previa = personalizada.ler(sessao, _caminho(arquivo), _layout_salvo(sessao, layout),
                                   atualizar_existentes)
        for linha in previa.linhas:
            if linha.linha in pular:
                linha.importar = False
        resultado = personalizada.importar(sessao, previa, atualizar_existentes)
        return {"resumo": resultado.resumo, "importacao_id": resultado.importacao_id,
                "falhas": [_problema(p) for p in resultado.falhas]}


# -------------------------------------------------------- cadastros de apoio
@_ferramenta(GRAVA)
def criar_tipo_de_honorario(
    nome: str,
    grupo_id: int,
    periodicidade: str = "MENSAL",
    exige_competencia: str = "MES",
    texto_mensagem: str = "",
    texto_mensagem_empresa: str = "",
    texto_item_lista: str = "",
) -> dict:
    """Novo tipo de honorário. periodicidade: MENSAL, TRIMESTRAL, ANUAL ou AVULSO;
    exige_competencia: MES, ANO ou NENHUMA. Textos vazios usam o nome do tipo."""
    from .dominio.enums import ExigeCompetencia

    with _sessao(gravar=True) as sessao:
        tipo = tipos.salvar_tipo(
            sessao, None, nome, grupo_id, _enum(Periodicidade, periodicidade),
            _enum(ExigeCompetencia, exige_competencia), texto_mensagem or nome,
            texto_mensagem_empresa or texto_mensagem or nome, texto_item_lista or nome)
        return {"id": tipo.id, "nome": tipo.nome}


@_ferramenta(GRAVA)
def criar_grupo_de_honorario(nome: str) -> dict:
    """Novo grupo de honorários (as seções das telas de cobrança)."""
    with _sessao(gravar=True) as sessao:
        grupo = tipos.criar_grupo(sessao, nome)
        return {"id": grupo.id, "nome": grupo.nome}


@_ferramenta(GRAVA)
def criar_forma_de_envio(nome: str) -> dict:
    """Nova forma de envio do boleto (ex.: e-mail, WhatsApp, portal)."""
    with _sessao(gravar=True) as sessao:
        forma = tipos.criar_forma(sessao, nome)
        return {"id": forma.id, "nome": forma.nome}


@_ferramenta(GRAVA)
def gerar_honorarios_dos_contratos(
    competencia_ano: int, competencia_mes: int, cliente_id: int | None = None,
    gravar: bool = False,
) -> dict:
    """Gera os honorários de uma competência a partir dos contratos (como a tela de
    geração): suspensos, fora da vigência e já lançados ficam de fora com o motivo. Com
    ``gravar=false`` (padrão) só mostra o que geraria."""
    with _sessao(gravar=gravar) as sessao:
        itens = contratos.previa_geracao(
            sessao, Competencia(ano=competencia_ano, mes=competencia_mes), cliente_id)
        criados = contratos.gerar(sessao, itens) if gravar else []
        return {
            "gravado": gravar, "gerados": len(criados),
            "itens": [{"contrato_id": i.contrato.id, "cliente": i.cliente.nome,
                       "tipo": i.contrato.tipo.nome, "valor": _plano(i.valor),
                       "vencimento": _plano(i.vencimento), "gera": i.gerar,
                       "impedimento": i.impedimento or None} for i in itens],
        }


@_ferramenta(GRAVA)
def registrar_acordo_de_valor(honorario_ids: list[int], valido_ate: str,
                              data_base: str | None = None) -> dict:
    """Congela o valor (com os juros até ``data_base``, hoje se vazia) até ``valido_ate``:
    pagando até lá, vale o valor do acordo."""
    from .servicos import acordos

    with _sessao(gravar=True) as sessao:
        criados = acordos.registrar(sessao, [_honorario(sessao, i) for i in honorario_ids],
                                    _data(data_base) or hoje(), _data(valido_ate))
        return {"acordos": [{"honorario_id": a.honorario_id, "total": _plano(a.total),
                             "valido_ate": _plano(a.valido_ate)} for a in criados]}


@_ferramenta(GRAVA)
def renegociar(
    cliente_id: int,
    honorario_ids: list[int],
    valor_acordado: str,
    parcelas: int,
    primeiro_vencimento: str,
    tipo_id: int,
    desconto: str = "0",
    observacoes: str = "",
) -> dict:
    """Renegocia honorários em aberto do cliente em parcelas (os originais ficam como
    renegociados e as parcelas entram como honorários novos do ``tipo_id``)."""
    from .servicos import renegociacoes

    with _sessao(gravar=True) as sessao:
        referencia = hoje()
        acordo = renegociacoes.criar(
            sessao, _cliente(sessao, cliente_id),
            [_linha(sessao, i, referencia) for i in honorario_ids],
            _valor(valor_acordado), parcelas, _data(primeiro_vencimento), tipo_id,
            _valor(desconto), observacoes)
        return {"renegociacao_id": acordo.id, "valor_acordado": _plano(acordo.valor_acordado),
                "parcelas": acordo.parcelas}


@_ferramenta(GRAVA)
def gerar_backup(motivo: str = "pedido pelo Claude") -> str:
    """Faz um backup completo do banco agora (o mesmo do botão Gerar backup)."""
    from .servicos import backup

    with _sessao():   # banco configurado e erro de negócio como mensagem
        return str(backup.gerar(motivo))


@_ferramenta(LEITURA)
def listar_configuracoes() -> dict:
    """Todas as configurações do sistema (chave → valor)."""
    from .infra.modelos import Configuracao

    with _sessao() as sessao:
        return {c.chave: c.valor for c in sessao.scalars(
            select(Configuracao).order_by(Configuracao.chave))}


# -------------------------------------------------------- acesso direto ao banco
def _executar_sql(sessao: Session, sql: str, parametros: dict | None):
    from sqlalchemy.exc import DBAPIError, StatementError

    if not (sql or "").strip():
        raise ValueError("Informe o comando SQL.")
    try:
        return sessao.execute(text(sql), parametros or {})
    except (DBAPIError, StatementError) as falha:
        motivo = getattr(falha, "orig", None) or falha
        raise ValueError(f"O banco recusou o comando: {motivo}") from falha


def _linhas_sql(resultado, limite: int) -> dict:
    colunas = list(resultado.keys())
    linhas = resultado.fetchmany(limite + 1)
    return {"colunas": colunas,
            "linhas": [[_plano(v) if not isinstance(v, bytes) else f"<{len(v)} bytes>"
                        for v in linha] for linha in linhas[:limite]],
            "mais_linhas": len(linhas) > limite}


@_ferramenta(LEITURA)
def esquema_do_banco(tabela: str | None = None) -> dict:
    """As tabelas do banco com as colunas (tipo, nulo, chave) e as ligações entre elas.
    Use antes de consultar_banco/executar_no_banco."""
    from sqlalchemy import inspect

    with _sessao() as sessao:
        inspetor = inspect(sessao.get_bind())
        nomes = [tabela] if tabela else sorted(inspetor.get_table_names())
        resultado = {}
        for nome in nomes:
            chave = set(inspetor.get_pk_constraint(nome).get("constrained_columns") or [])
            resultado[nome] = {
                "colunas": [{"nome": c["name"], "tipo": str(c["type"]),
                             "nulo": c["nullable"], "chave": c["name"] in chave}
                            for c in inspetor.get_columns(nome)],
                "ligacoes": [{"colunas": f["constrained_columns"],
                              "tabela": f["referred_table"], "para": f["referred_columns"]}
                             for f in inspetor.get_foreign_keys(nome)],
            }
        return {"banco": sessao.get_bind().dialect.name, "tabelas": resultado}


@_ferramenta(LEITURA)
def consultar_banco(sql: str, parametros: dict | None = None, limite: int = 500) -> dict:
    """Roda uma consulta SQL (SELECT) direto no banco e devolve as linhas. Nada é
    gravado: a transação é sempre desfeita. Use :nome no SQL e ``parametros`` para
    valores (ex.: "select * from cliente where codigo = :codigo")."""
    with _sessao() as sessao:
        if sessao.get_bind().dialect.name == "postgresql":
            sessao.execute(text("SET TRANSACTION READ ONLY"))
        resultado = _executar_sql(sessao, sql, parametros)
        if not resultado.returns_rows:
            raise ValueError("Esse comando não devolve linhas: para gravar, use "
                             "executar_no_banco.")
        dados = _linhas_sql(resultado, limite)
        sessao.rollback()
        return dados


@_ferramenta(APAGA)
def executar_no_banco(sql: str, parametros: dict | None = None) -> dict:
    """Roda um comando SQL que grava (INSERT, UPDATE, DELETE, ALTER…) direto no banco e
    confirma. Passa por fora das regras do sistema (juros, validações, histórico):
    prefira as ferramentas específicas e use isto para o que elas não cobrem. Um comando
    por vez; fica anotado no log de alterações. Confirme com o usuário antes."""
    from .infra.modelos import LogAlteracao

    with _sessao(gravar=True) as sessao:
        resultado = _executar_sql(sessao, sql, parametros)
        retorno = _linhas_sql(resultado, 500) if resultado.returns_rows else None
        afetadas = resultado.rowcount
        sessao.add(LogAlteracao(entidade="banco", acao="SQL", resumo=sql.strip()[:4000],
                                depois=json.dumps(_plano(parametros or {}), ensure_ascii=False)))
        return {"linhas_afetadas": afetadas, "retorno": retorno}


def main() -> None:
    # stdout é o canal do protocolo: nada de print; erros vão para o stderr
    try:
        mcp.run()
    except KeyboardInterrupt:
        pass
    except Exception as falha:  # noqa: BLE001
        print(f"Servidor MCP parou: {falha}", file=sys.stderr)
        raise SystemExit(1) from falha


if __name__ == "__main__":
    main()
